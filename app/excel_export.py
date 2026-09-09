"""Genera el Excel del line-up reproduciendo el formato del archivo original.

Dos variantes:
  * interno  -> incluye LOCAL, PRINCIPAL y OTRAS AGENCIAS. Nombre completo.
  * clientes -> sin LOCAL / PRINCIPAL / OTRAS AGENCIAS. Nombre sin la
                primera palabra del prefijo.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import settings
from .dates import parse_date

FONT = "Courier New"
PORT_FILL = PatternFill("solid", fgColor="FF009874")
HEADER_FILL = PatternFill("solid", fgColor="C6E0B4")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFFFF")
THIN = Side(style="thin", color="FFBFBFBF")

# (encabezado, atributo del VesselCall, ancho de columna, es_fecha)
BASE_COLS = [
    ("__TERMINAL__", "vessel_name", 26.9, False),
    ("Vessel type", "vessel_type", 21.3, False),
    ("IMO", "imo", 16.0, False),
    ("ETA", "eta", 17.6, True),
    ("ETB", "etb", 18.6, True),
    ("ETC", "etc", 18.4, True),
    ("Operation", "operation", 14.0, False),
    ("Quantity", "quantity", 15.7, False),
    ("GRADE", "grade", 16.7, False),
    ("SHIPPER", "shipper", 12.5, False),
    ("DESTINATION", "destination", 19.4, False),
]
INTERNAL_EXTRA = [
    ("LOCAL", "local_agent", 14.9, False),
    ("PRINCIPAL", "principal_name", 14.0, False),
    ("OTRAS AGENCIAS", "__extra__", 22.0, False),
]


def _write_value(cell, value: str, is_date: bool):
    if is_date:
        d = parse_date(value)
        if d:
            cell.value = d
            cell.number_format = "dd/mm/yy;@"
            return
    text = "" if value is None else str(value)
    if text.isdigit() and len(text) <= 15:
        cell.value = int(text)
    else:
        cell.value = text


def build_lineup_xlsx(lineup, terminals, calls_by_terminal, *, internal: bool) -> tuple[io.BytesIO, str]:
    cols = BASE_COLS + (INTERNAL_EXTRA if internal else [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja2"
    ws.sheet_view.showGridLines = False

    first_col = 2  # arranca en B (A es un margen angosto)
    last_col = first_col + len(cols) - 1

    ws.column_dimensions[get_column_letter(1)].width = 7.2
    for i, (_, _, width, _) in enumerate(cols):
        ws.column_dimensions[get_column_letter(first_col + i)].width = width

    # Fila 2: nombre del puerto
    ws.row_dimensions[2].height = 15.6
    ws.merge_cells(start_row=2, start_column=first_col, end_row=2, end_column=min(last_col, first_col + 10))
    top = ws.cell(row=2, column=first_col, value=(lineup.port_name or settings.port_name) + " ")
    top.font = Font(name="Calibri", size=12, bold=True, color="FFFFFFFF")
    top.fill = PORT_FILL
    top.alignment = Alignment(horizontal="center", vertical="center")

    row = 4
    for term in terminals:
        term_calls = calls_by_terminal.get(term.id, [])

        # Encabezado del bloque
        for i, (head, _, _, is_date) in enumerate(cols):
            c = ws.cell(row=row, column=first_col + i)
            c.value = term.code if head == "__TERMINAL__" else head
            c.font = Font(name=FONT, size=10, bold=True)
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.row_dimensions[row].height = 15.0
        row += 1

        for call in term_calls:
            for i, (head, attr, _, is_date) in enumerate(cols):
                c = ws.cell(row=row, column=first_col + i)
                if attr == "__extra__":
                    val = ", ".join(cl.name for cl in _extra_clients(call))
                else:
                    val = getattr(call, attr, "")
                _write_value(c, val, is_date)
                c.font = Font(name=FONT, size=10)
                c.fill = WHITE_FILL
                c.alignment = Alignment(horizontal="left", vertical="center")
            row += 1

        row += 1  # fila en blanco entre bloques

    # Windguru
    row += 1
    wg = ws.cell(row=row, column=first_col, value=settings.windguru_label)
    wg.hyperlink = settings.windguru_url
    wg.font = Font(name="Calibri", size=11, color="FF0563C1", underline="single")

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    date_tag = (lineup.lineup_date or "").replace("/", ".")
    prefix = settings.internal_filename_prefix if internal else settings.client_filename_prefix
    filename = f"{prefix} {date_tag}.xlsx".strip()
    return bio, filename


def _extra_clients(call):
    return [link.client for link in getattr(call, "extra_agencies", []) if link.client]
