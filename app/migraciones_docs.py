"""Generacion de notas de Migraciones (embarco/desembarco), acta de
reconduccion y shore pass (ficha individual ANEXO VI).

En vez de reconstruir el formato a mano, se parte de los .docx/.xlsx
reales que usa la agencia (guardados -ya sin los datos de ejemplo- en
app/migraciones_templates/) y se completan por posicion de run/celda.
Asi el header con el logo, la tipografia, los margenes y el tamano de
pagina salen identicos a los que ya se usaban en papel.

One-shot / stateless -- igual que liquidaciones.py, no persiste nada."""
from __future__ import annotations

import copy
import io
from dataclasses import dataclass
from datetime import date

import docx
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage

from .config import BASE_DIR

TEMPLATES_DIR = BASE_DIR / "app" / "migraciones_templates"

_MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]


def _fecha_larga(d: date) -> str:
    return f"{d.day} de {_MESES[d.month - 1]} de {d.year}"


@dataclass
class Tripulante:
    nombre: str = ""
    nacionalidad: str = ""
    fecha_nacimiento: str = ""
    lugar_nacimiento: str = ""
    pasaporte: str = ""
    rango: str = ""
    vuelo: str = ""
    fecha_salida: str = ""
    hora_salida: str = ""
    expedida_por: str = ""


# --------------------------------------------------------------------- #
# Nota de embarco / desembarco
# --------------------------------------------------------------------- #

def _insertar_parrafo_despues(doc, ref_element, texto: str) -> Paragraph:
    """Crea un parrafo nuevo (Courier New 11, igual que el resto de la
    nota) y lo reubica en el arbol XML justo despues de `ref_element`."""
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.font.name = "Courier New"
    run.font.size = Pt(11)
    ref_element.addnext(p._p)
    return p


def build_nota_migraciones(
    *, tipo: str, fecha: date, bandera: str, buque: str, muelle: str,
    tripulantes: list[Tripulante],
) -> io.BytesIO:
    """tipo: 'EMBARCO' o 'DESEMBARCO'. Parte de migraciones_templates/nota.docx
    (el .doc real de la agencia, sanitizado) y completa por posicion de run."""
    es_desembarco = tipo.upper() == "DESEMBARCO"

    doc = docx.Document(str(TEMPLATES_DIR / "nota.docx"))
    paras = doc.paragraphs

    # P7: "BAHIA BLANCA, <dia> de <MES> de <ano>."
    p7 = paras[7]
    p7.runs[2].text = str(fecha.day)
    p7.runs[4].text = _MESES[fecha.month - 1]
    p7.runs[5].text = " de "
    p7.runs[6].text = str(fecha.year)

    # P14: parrafo de autorizacion -- embarco/desembarco, bandera, buque, muelle
    p14 = paras[14]
    p14.runs[2].text = "des" if es_desembarco else ""
    p14.runs[12].text = bandera
    p14.runs[14].text = buque
    if muelle.strip():
        p14.runs[15].text = '", el cual se encuentra'
        p14.runs[16].text = " en "
        p14.runs[17].text = f"Muelle {muelle}"
        p14.runs[18].text = " "
    else:
        p14.runs[15].text = '"'
        p14.runs[16].text = ""
        p14.runs[17].text = ""
        p14.runs[18].text = ""

    # P15..P20: bloque de un tripulante (blanco, nombre, nacionalidad,
    # fecha nac., pasaporte, rango) -- se clona una vez por cada tripulante
    # de mas, y se llena en orden.
    pristino = [paras[i]._p for i in range(15, 21)]
    anchor = pristino[-1]
    bloques = [pristino]
    for _ in tripulantes[1:]:
        clones = [copy.deepcopy(e) for e in pristino]
        prev = anchor
        for c in clones:
            prev.addnext(c)
            prev = c
        anchor = clones[-1]
        bloques.append(clones)

    for bloque_elems, trip in zip(bloques, tripulantes):
        _, p_nombre, p_nac, p_fecha, p_pasaporte, p_rango = (Paragraph(e, doc) for e in bloque_elems)
        p_nombre.runs[0].text = trip.nombre.upper()
        p_nac.runs[1].text = trip.nacionalidad
        p_fecha.runs[1].text = trip.fecha_nacimiento
        p_pasaporte.runs[1].text = trip.pasaporte
        p_rango.runs[1].text = trip.rango
        if es_desembarco:
            texto = f"Salida del pais: {trip.vuelo} – {trip.fecha_salida} a las {trip.hora_salida} hs"
            _insertar_parrafo_despues(doc, bloque_elems[-1], texto)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# --------------------------------------------------------------------- #
# Acta de reconduccion
# --------------------------------------------------------------------- #

def build_acta_reconduccion(
    *, fecha: date, nacionalidad: str, nombre: str, pasaporte: str,
    bandera: str, buque: str, vuelo: str, fecha_salida: str, hora_salida: str,
) -> io.BytesIO:
    doc = docx.Document(str(TEMPLATES_DIR / "acta_reconduccion.docx"))
    paras = doc.paragraphs

    p7 = paras[7]
    p7.runs[4].text = f" {nacionalidad.upper()}"
    p7.runs[6].text = nombre.upper()
    p7.runs[9].text = f" {pasaporte}"
    p7.runs[12].text = bandera.upper()
    p7.runs[14].text = buque.upper()
    p7.runs[22].text = vuelo
    p7.runs[24].text = fecha_salida
    p7.runs[25].text = ""
    p7.runs[26].text = ""
    p7.runs[29].text = hora_salida
    p7.runs[30].text = ""
    p7.runs[31].text = ""

    p9 = paras[9]
    p9.runs[1].text = str(fecha.day)
    p9.runs[3].text = f"{_MESES[fecha.month - 1]} "
    p9.runs[4].text = "de "
    p9.runs[5].text = str(fecha.year)
    p9.runs[6].text = ""

    # Numero de expediente DNM -- por ahora siempre el mismo (258369)
    paras[16].runs[2].text = "  258369"

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# --------------------------------------------------------------------- #
# Shore pass (Ficha individual tripulante - ANEXO VI)
# --------------------------------------------------------------------- #

def _recortar_cuadro(foto_bytes: bytes, target_w: int, target_h: int) -> io.BytesIO:
    """Recorta (centrado, sin deformar) y escala la foto para que llene
    exactamente el cuadro del ANEXO VI, en vez de dejarla mas chica con
    margenes si no coincide la proporcion del original."""
    img = PILImage.open(io.BytesIO(foto_bytes)).convert("RGB")
    ratio_caja = target_w / target_h
    ancho, alto = img.size
    ratio_foto = ancho / alto
    if ratio_foto > ratio_caja:
        nuevo_ancho = int(alto * ratio_caja)
        x0 = (ancho - nuevo_ancho) // 2
        img = img.crop((x0, 0, x0 + nuevo_ancho, alto))
    else:
        nuevo_alto = int(ancho / ratio_caja)
        y0 = (alto - nuevo_alto) // 2
        img = img.crop((0, y0, ancho, y0 + nuevo_alto))
    img = img.resize((target_w, target_h), PILImage.LANCZOS)
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def build_shore_pass(
    *, buque: str, empresa_signataria: str, lugar_fecha: str, trip: Tripulante,
    foto_bytes: bytes | None,
) -> io.BytesIO:
    """Parte del modelo real (migraciones_templates/shore_pass.xlsx, ya con
    el logo y todas las cajas/bordes del ANEXO VI) y solo completa los
    valores en las celdas que corresponden."""
    wb = load_workbook(str(TEMPLATES_DIR / "shore_pass.xlsx"))
    ws = wb.active

    ws["J8"] = buque
    ws["O11"] = empresa_signataria
    ws["U14"] = trip.nombre.upper()
    ws["S17"] = f"{trip.fecha_nacimiento}  {trip.lugar_nacimiento}".strip()
    ws["L20"] = trip.nacionalidad
    ws["M24"] = trip.rango
    ws["AJ24"] = trip.pasaporte
    ws["L27"] = trip.expedida_por or trip.nacionalidad
    ws["V37"] = lugar_fecha

    if foto_bytes:
        target_w, target_h = 190, 165  # caja D33:L40 del modelo
        recortada = _recortar_cuadro(foto_bytes, target_w, target_h)
        foto = XLImage(recortada)
        foto.width, foto.height = target_w, target_h
        marker = AnchorMarker(col=3, colOff=pixels_to_EMU(6), row=32, rowOff=pixels_to_EMU(4))
        foto.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(target_w), pixels_to_EMU(target_h)))
        ws.add_image(foto)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
