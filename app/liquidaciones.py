"""Clasificacion automatica de habilitaciones de Servicios Extraordinarios
(Liquidaciones Aduana). Modulo puro (sin FastAPI ni DB): matchea el Excel
de "Consulta de liquidaciones" de ARCA contra un ZIP de PDFs de solicitud
de habilitacion (OM 1077-C) por numero de habilitacion, completa BUQUE y
TIPO DE HABILITACION, y genera un Excel de salida con resumenes.

One-shot / stateless -- no persiste nada, no se integra con escalas."""
import io
import re
import zipfile
from dataclasses import dataclass, field

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RE_TIPO = re.compile(r"Tipo de Operacion\s+(.+?)(?:\n|$)")
RE_BUQUE = re.compile(r"Nombre del Medio Transportador\s*(.*?)(?:\s+Via\b|\n)")
RE_LUGAR = re.compile(r"Lugar en que se realizara\s+(.*?)(?:\s+Zona\b|\n)")
RE_REF = re.compile(r"\d{5}SSEE\d{6}[A-Z]")

# Terminos conocidos que aparecen en "Nombre del Medio Transportador" pero
# no son buques -- se muestran igual en el resumen, solo marcados aparte.
NO_ES_BUQUE = {"TERMINAL PATAGONIA NORTE", "VARIOS"}

# Alias de nombres de buque que no estan normalizados entre el PDF y otros
# lugares -- punto de extension natural cuando se integre con escalas.
ALIAS_BUQUE = {"T PRIME": "TP PRIME"}

# Codigo de 4 letras -> descripcion completa, solo para mostrar compacto.
# Lista abierta: pueden aparecer codigos nuevos, no se valida contra esto.
CATALOGO_TIPOS = {
    "CARB": "Carga de Buques",
    "META": "Mediciones en Tanques Fiscales",
    "CUST": "Custodia",
    "PERM": "Permanencias",
    "FEVA": "Formalizacion de Entrada Via Acuatica",
    "DEBU": "Descarga de Buques",
    "PART": "Solicitud Particular OM 1247",
    "RANP": "Ranchos y Provisiones",
    "RANC": "Rancho de Combustibles",
    "PREM": "Presentacion SIM",
}

BLUE = "FF1B5FA8"
TOTAL_FILL = "FFDDEBF7"
WHITE_BOLD = Font(bold=True, color="FFFFFFFF")
BOLD = Font(bold=True)


def parse_pdf(fileobj) -> dict:
    """Extrae tipo de operacion, buque y lugar de una solicitud OM 1077-C."""
    with pdfplumber.open(fileobj) as pdf:
        texto = "\n".join((pg.extract_text() or "") for pg in pdf.pages)

    m_tipo = RE_TIPO.search(texto)
    m_buque = RE_BUQUE.search(texto)
    m_lugar = RE_LUGAR.search(texto)

    return {
        "tipo": m_tipo.group(1).strip() if m_tipo else "",
        "buque": m_buque.group(1).strip() if m_buque else "",
        "lugar": m_lugar.group(1).strip() if m_lugar else "",
    }


def ref_desde_nombre(filename: str) -> str | None:
    """'26003SSEE018282N (1).pdf' -> '26003SSEE018282N'"""
    m = RE_REF.search(filename.upper())
    return m.group(0) if m else None


def normalizar_buque(nombre: str) -> str:
    return ALIAS_BUQUE.get(nombre, nombre)


CARPETA_SIN_CLASIFICAR = "SIN CLASIFICAR"
_RE_CHARS_INVALIDOS = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def nombre_carpeta(buque: str) -> str:
    """Nombre de carpeta seguro para Windows a partir del buque del PDF.
    Vacio -> SIN CLASIFICAR (el PDF no tenia el campo o no se pudo leer)."""
    limpio = _RE_CHARS_INVALIDOS.sub(" ", buque or "")
    limpio = re.sub(r"\s+", " ", limpio).strip(" .")
    if not limpio or limpio.upper() == CARPETA_SIN_CLASIFICAR:
        return CARPETA_SIN_CLASIFICAR
    return limpio


def _monto_a_float(v) -> float:
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


def _codigo_tipo(tipo_completo: str) -> str:
    """'CARB - Carga de Buques' -> 'CARB' (el codigo es lo primero antes de
    ' - ', o el string entero si no tiene ese separador)."""
    return (tipo_completo or "").split(" - ", 1)[0].strip()


@dataclass
class ResultadoLiquidaciones:
    total_habilitaciones: int = 0
    total_monto: float = 0.0
    filas_sin_pdf: list[str] = field(default_factory=list)
    pdfs_sin_fila: list[str] = field(default_factory=list)
    duplicados: list[str] = field(default_factory=list)
    codigos_desconocidos: list[str] = field(default_factory=list)
    resumen_buque: list[dict] = field(default_factory=list)
    resumen_tipo: list[dict] = field(default_factory=list)
    aliases_aplicados: dict = field(default_factory=dict)
    excel_bytes: bytes = b""
    # ZIP con los PDFs ordenados en una carpeta por buque + SIN CLASIFICAR
    zip_bytes: bytes = b""
    carpetas: list[dict] = field(default_factory=list)


def _encontrar_hoja_datos(wb):
    for ws in wb.worksheets:
        if str(ws["A1"].value or "").strip().lower().startswith("liquidaci"):
            return ws
    return None


def _encontrar_columna_referencia(ws, fila_inicio: int, fila_fin: int) -> int:
    """Columna L no tiene encabezado en el export de ARCA. Se busca por
    contenido (primera columna cuyas celdas matcheen RE_REF); si no se
    encuentra ninguna, cae a la posicion fija (12, 1-based)."""
    max_col = ws.max_column
    for col in range(1, max_col + 1):
        matches = 0
        muestras = 0
        for row in range(fila_inicio, min(fila_fin, fila_inicio + 20) + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            muestras += 1
            if RE_REF.search(str(v).upper()):
                matches += 1
        if muestras and matches == muestras:
            return col
    return 12


def procesar(excel_bytes: bytes, zip_bytes: bytes) -> ResultadoLiquidaciones:
    resultado = ResultadoLiquidaciones()

    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    ws = _encontrar_hoja_datos(wb)
    if ws is None:
        raise ValueError('No se encontro ninguna hoja con "Liquidación" en A1.')

    fila_inicio = 2
    fila_fin = ws.max_row
    while fila_fin >= fila_inicio and all(
        ws.cell(row=fila_fin, column=c).value in (None, "") for c in range(1, ws.max_column + 1)
    ):
        fila_fin -= 1

    col_ref = _encontrar_columna_referencia(ws, fila_inicio, fila_fin)
    col_monto = 11  # K, "Monto Documento Referencia"

    # columnas N (14) y O (15) -- crearlas si el excel viniera sin ellas
    if ws.max_column < 15:
        ws.cell(row=1, column=14, value="BUQUE")
        ws.cell(row=1, column=15, value="TIPO DE HABILITACION")
    else:
        if not ws.cell(row=1, column=14).value:
            ws.cell(row=1, column=14, value="BUQUE")
        if not ws.cell(row=1, column=15).value:
            ws.cell(row=1, column=15, value="TIPO DE HABILITACION")

    # --- parsear todos los PDFs del zip ---
    pdfs_por_ref: dict[str, dict] = {}
    nombres_vistos: dict[str, str] = {}
    # Cada PDF se copia al ZIP de salida en la carpeta de su buque. Lo que no
    # se puede leer, no tiene numero de habilitacion en el nombre o no trae
    # buque va a SIN CLASIFICAR, para que nada se pierda ni quede mal ubicado.
    carpetas: dict[str, int] = {}
    nombres_usados: set[str] = set()
    out_bio = io.BytesIO()

    def _guardar_en_zip(zout, base: str, contenido: bytes, buque: str):
        carpeta = nombre_carpeta(buque)
        destino = f"{carpeta}/{base}"
        n = 1
        while destino.lower() in nombres_usados:
            n += 1
            stem, dot, ext = base.rpartition(".")
            destino = f"{carpeta}/{stem} ({n}).{ext}" if dot else f"{carpeta}/{base} ({n})"
        nombres_usados.add(destino.lower())
        zout.writestr(destino, contenido)
        carpetas[carpeta] = carpetas.get(carpeta, 0) + 1

    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf, zipfile.ZipFile(out_bio, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zf.infolist():
            name = info.filename
            base = name.rsplit("/", 1)[-1]
            if info.is_dir() or not base.lower().endswith(".pdf") or name.startswith("__MACOSX"):
                continue
            contenido = zf.read(info)
            ref = ref_desde_nombre(base)
            try:
                datos = parse_pdf(io.BytesIO(contenido))
            except Exception:
                datos = {"tipo": "", "buque": "", "lugar": ""}
            buque_pdf = normalizar_buque(datos["buque"])

            if not ref:
                _guardar_en_zip(zout, base, contenido, "")
                continue
            if ref in pdfs_por_ref:
                resultado.duplicados.append(base)
                _guardar_en_zip(zout, base, contenido, buque_pdf)
                continue
            nombres_vistos[ref] = base
            pdfs_por_ref[ref] = datos
            _guardar_en_zip(zout, base, contenido, buque_pdf)
    resultado.zip_bytes = out_bio.getvalue()
    resultado.carpetas = sorted(
        [{"nombre": n, "cantidad": c, "sin_clasificar": n == CARPETA_SIN_CLASIFICAR} for n, c in carpetas.items()],
        key=lambda x: (x["sin_clasificar"], x["nombre"]),
    )

    # --- matchear filas del excel contra los pdfs ---
    refs_usadas: set[str] = set()
    por_buque: dict[str, dict] = {}
    por_tipo: dict[str, dict] = {}
    codigos_vistos: set[str] = set()

    for row in range(fila_inicio, fila_fin + 1):
        raw_ref = ws.cell(row=row, column=col_ref).value
        if raw_ref in (None, ""):
            continue
        ref = ref_desde_nombre(str(raw_ref))
        monto = _monto_a_float(ws.cell(row=row, column=col_monto).value)
        # reescribe el monto como numero real, con formato
        celda_monto = ws.cell(row=row, column=col_monto, value=monto)
        celda_monto.number_format = "#,##0.00"

        datos = pdfs_por_ref.get(ref) if ref else None
        if not datos:
            liq = ws.cell(row=row, column=1).value
            resultado.filas_sin_pdf.append(str(liq or f"fila {row}"))
            continue
        refs_usadas.add(ref)

        buque_original = datos["buque"]
        buque = normalizar_buque(buque_original)
        if buque != buque_original:
            resultado.aliases_aplicados[buque_original] = buque
        tipo_completo = datos["tipo"]
        codigo = _codigo_tipo(tipo_completo)

        ws.cell(row=row, column=14, value=buque)
        ws.cell(row=row, column=15, value=tipo_completo)

        resultado.total_habilitaciones += 1
        resultado.total_monto += monto

        entry = por_buque.setdefault(buque, {"cantidad": 0, "monto": 0.0, "tipos": {}, "no_es_buque": buque in NO_ES_BUQUE})
        entry["cantidad"] += 1
        entry["monto"] += monto
        entry["tipos"][codigo] = entry["tipos"].get(codigo, 0) + 1

        tentry = por_tipo.setdefault(codigo, {"descripcion": tipo_completo, "cantidad": 0, "monto": 0.0})
        tentry["cantidad"] += 1
        tentry["monto"] += monto

        if codigo and codigo not in CATALOGO_TIPOS:
            codigos_vistos.add(codigo)

    for ref, base in nombres_vistos.items():
        if ref not in refs_usadas:
            resultado.pdfs_sin_fila.append(base)

    resultado.codigos_desconocidos = sorted(codigos_vistos)

    total = resultado.total_monto or 1.0
    resultado.resumen_buque = sorted(
        [
            {
                "buque": b,
                "cantidad": e["cantidad"],
                "monto": e["monto"],
                "pct": e["monto"] / total,
                "tipos": ", ".join(f"{k}:{v}" for k, v in sorted(e["tipos"].items(), key=lambda x: -x[1])),
                "no_es_buque": e["no_es_buque"],
            }
            for b, e in por_buque.items()
        ],
        key=lambda x: -x["monto"],
    )
    resultado.resumen_tipo = sorted(
        [{"codigo": k, "descripcion": t["descripcion"], "cantidad": t["cantidad"], "monto": t["monto"]} for k, t in por_tipo.items()],
        key=lambda x: -x["monto"],
    )

    _armar_hoja_datos(ws)
    _insertar_hoja_resumen(wb, ws, resultado)
    _insertar_hoja_por_tipo(wb, resultado)

    bio = io.BytesIO()
    wb.save(bio)
    resultado.excel_bytes = bio.getvalue()
    return resultado


def _armar_hoja_datos(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["N"].width = 24
    ws.column_dimensions["O"].width = 30


def _fill(color):
    return PatternFill("solid", fgColor=color)


def _insertar_hoja_resumen(wb, ws_datos, resultado: ResultadoLiquidaciones):
    idx = wb.worksheets.index(ws_datos) + 2  # justo despues de la hoja de datos
    ws = wb.create_sheet("Resumen", idx - 1)
    headers = ["BUQUE / MEDIO", "Habilitaciones", "Monto total (PES)", "% del total", "Tipos de habilitacion"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = _fill(BLUE)
    row = 2
    for r in resultado.resumen_buque:
        etiqueta = r["buque"] + (" (no es buque)" if r["no_es_buque"] else "")
        ws.cell(row=row, column=1, value=etiqueta)
        ws.cell(row=row, column=2, value=r["cantidad"])
        c = ws.cell(row=row, column=3, value=r["monto"])
        c.number_format = "#,##0.00"
        c = ws.cell(row=row, column=4, value=r["pct"])
        c.number_format = "0.0%"
        ws.cell(row=row, column=5, value=r["tipos"])
        row += 1
    for i in range(1, 6):
        c = ws.cell(row=row, column=i)
        c.font = BOLD
        c.fill = _fill(TOTAL_FILL)
    ws.cell(row=row, column=1, value="TOTAL")
    ws.cell(row=row, column=2, value=resultado.total_habilitaciones)
    c = ws.cell(row=row, column=3, value=resultado.total_monto)
    c.number_format = "#,##0.00"
    ws.cell(row=row, column=4, value=1)
    ws.cell(row=row, column=4).number_format = "0.0%"
    for i, w in enumerate([32, 16, 20, 12, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _insertar_hoja_por_tipo(wb, resultado: ResultadoLiquidaciones):
    ws = wb.create_sheet("Por tipo", 2)
    headers = ["TIPO DE HABILITACION", "Habilitaciones", "Monto total (PES)"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = WHITE_BOLD
        c.fill = _fill(BLUE)
    row = 2
    for r in resultado.resumen_tipo:
        etiqueta = f"{r['codigo']} - {CATALOGO_TIPOS.get(r['codigo'], r['descripcion'])}"
        ws.cell(row=row, column=1, value=etiqueta)
        ws.cell(row=row, column=2, value=r["cantidad"])
        c = ws.cell(row=row, column=3, value=r["monto"])
        c.number_format = "#,##0.00"
        row += 1
    for i in range(1, 4):
        c = ws.cell(row=row, column=i)
        c.font = BOLD
        c.fill = _fill(TOTAL_FILL)
    ws.cell(row=row, column=1, value="TOTAL")
    ws.cell(row=row, column=2, value=resultado.total_habilitaciones)
    c = ws.cell(row=row, column=3, value=resultado.total_monto)
    c.number_format = "#,##0.00"
    for i, w in enumerate([40, 16, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
