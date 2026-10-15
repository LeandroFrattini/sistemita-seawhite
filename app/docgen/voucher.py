"""Voucher (Service Certificate): conceptos por formato, lineas que salen segun lo tildado y los documentos.

Dos formatos: el general de Sea White (Excel, REG-AM-36, para todas las agencias) y el de OCEANWAY (Word).
Todos los conceptos vienen tildados de entrada y se destildan los que no corresponden. Hay tres tipos:
- simple: una linea.
- group: CUSTOM OVERTIME, con sub-conceptos (Entrance, Loading, Permanence, Bunkering, ...) que se tildan por separado.
- dir: migraciones y taxis, a veces solo IN, a veces solo OUT. Solo las MIGRACIONES pueden ir en boya ("AT BUOY 11",
  mas caro), cada direccion por separado; ningun otro concepto lleva "AT BUOY 11".
Ademas se pueden agregar conceptos a mano (los que ya se usaron quedan guardados para ofrecerlos de nuevo).
"""
from __future__ import annotations

import copy
import io
import re
from datetime import date
from pathlib import Path

import docx
from docx.oxml.ns import qn

from lxml import etree

from .xlsx_fill import M, Xlsx

TEMPLATES = Path(__file__).resolve().parent / "templates"
BOYA = " AT BUOY 11"
MAX_EXTRA = 30

EXCEL, WORD = "EXCEL", "WORD"

CUSTOM_SUBS = ["ENTRANCE", "LOADING", "PERMANENCE", "BUNKERING", "CLEARANCE", "STORE LIST ENDORSED"]


def _s(key, text, **kw):
    return {"key": key, "tipo": "simple", "text": text, **kw}


def _g(key, prefix, subs):
    return {"key": key, "tipo": "group", "prefix": prefix, "subs": subs}


def _d(key, base, boya=False):
    """Migraciones y taxis: IN y OUT. Solo las migraciones (boya=True) pueden ir 'AT BUOY 11' (en boya cuesta mas)."""
    return {"key": key, "tipo": "dir", "base": base, "boya": boya}


# Orden y redaccion tal cual los documentos de cada formato.
CATALOGO = {
    EXCEL: [
        _s("port_dues", "PORT DUES"),
        _s("channel_toll", "CHANNEL TOLL"),
        _s("tugboats", "TUGBOATS FOR BERTHING / UNBERTHING"),
        _s("pilot", "PILOT SERVICE IN / OUT"),
        _s("mooring", "MOORING / UNMOORING (LDS)"),
        _s("watchman", "COMPULSORY WATCHMAN SERVICES"),
        _s("head_tally", "HEAD TALLY CLERK EXPENSES", head_tally=True),
        _s("agency", "AGENCY EXPENSES"),
        _s("bank", "BANK COMISION / LAW 25413 / ETC"),
        _g("custom", "CUSTOM OVERTIME FOR ", CUSTOM_SUBS),
        _s("senasa", "EXPENSES FOR COMPULSORY SENASA INSPECTION (REGULATION 714/2010)"),
        _d("immigration", "IMMIGRATION TAX EXPENSES", boya=True),
        _d("taxi_immigration", "TAXI HIRED FOR IMMIGRATION OFFICER"),
        _s("free_pratique", "FREE PRATIQUE TAX"),
        _s("taxi_sanitary", "TAXI HIRED FOR SANITARY AUTHORITY"),
        _s("shore_gangway", "SHORE GANGWAY", shore_gangway=True),
    ],
    WORD: [
        _s("port_dues", "PORT DUES"),
        _s("channel_toll", "CHANNEL TOLL"),
        _s("tugboats", "TUGBOATS FOR BERTHING / UNBERTHING"),
        _s("pilot", "PILOT SERVICE IN / OUT"),
        _s("mooring", "MOORING / UNMOORING GANG SERVICE"),
        _s("head_tally", "HEAD TALLY CLERK EXPENSES", head_tally=True),
        _s("watchman", "WATCHMAN EXPENSES"),
        _s("osro", "OSRO CERTIFICATE FOR BUNKER"),
        _g("custom", "CUSTOM OVERTIME FOR ", CUSTOM_SUBS[:5]),
        _d("immigration", "IMMIGRATION EXPENSES", boya=True),
        _s("sanitary", "SANITARY FREE PRATIQUE"),
        _d("taxi_immigration", "TAXI HIRED FOR IMMIGRATION"),
        _s("taxi_sanitary", "TAXI HIRED FOR SANITARY AUTHORITIES"),
        _s("agency", "AGENCY EXPENSES"),
        _s("bank", "BANK COMISION / LAW 25413 / ETC"),
        _s("shore_gangway", "SHORE GANGWAY", shore_gangway=True),
    ],
}


def formato_para(agencia: str) -> str:
    """Solo OCEANWAY usa su propio formato (Word); el resto, el general de Sea White (Excel)."""
    return WORD if (agencia or "").strip().lower().startswith("oceanway") else EXCEL


def _estado_concepto(c: dict) -> dict:
    if c["tipo"] == "simple":
        return {"on": True, "buoy": False}
    if c["tipo"] == "group":
        return {"on": True, "subs": {s: {"on": True, "buoy": False} for s in c["subs"]}}
    return {"on": True, "in": True, "out": True, "in_buoy": False, "out_buoy": False}


def estado_inicial(formato: str) -> dict:
    """Todos los conceptos tildados; ustedes destildan los que no corresponden."""
    return {"c": {c["key"]: _estado_concepto(c) for c in CATALOGO[formato]}, "extra": []}


def normalizar(formato: str, raw) -> dict:
    """Estado valido a partir de lo que llega del navegador (solo claves conocidas, booleanos y textos acotados)."""
    raw = raw if isinstance(raw, dict) else {}
    rc = raw.get("c") if isinstance(raw.get("c"), dict) else {}
    out = estado_inicial(formato)
    for c in CATALOGO[formato]:
        src = rc.get(c["key"])
        if not isinstance(src, dict):
            continue
        dst = out["c"][c["key"]]
        dst["on"] = bool(src.get("on"))
        if c["tipo"] == "group":
            subs = src.get("subs") if isinstance(src.get("subs"), dict) else {}
            for name in c["subs"]:
                if isinstance(subs.get(name), dict):
                    dst["subs"][name] = {"on": bool(subs[name].get("on")), "buoy": False}
        elif c["tipo"] == "dir":
            for k in ("in", "out"):
                dst[k] = bool(src.get(k))
            if c.get("boya"):  # solo migraciones
                dst["in_buoy"], dst["out_buoy"] = bool(src.get("in_buoy")), bool(src.get("out_buoy"))
        # simple: sin boya (se descarta si venia de un voucher viejo)
    extra = []
    for e in (raw.get("extra") if isinstance(raw.get("extra"), list) else [])[:MAX_EXTRA]:
        if not isinstance(e, dict):
            continue
        t = " ".join(str(e.get("t") or "").split())[:200]
        if t:
            extra.append({"t": t.upper(), "on": bool(e.get("on", True)), "buoy": False})
    out["extra"] = extra
    return out


def lineas(formato: str, estado: dict) -> list[str]:
    """Las lineas que salen en el documento, en el orden del formato, mas los conceptos agregados a mano."""
    est = normalizar(formato, estado)
    out: list[str] = []
    for c in CATALOGO[formato]:
        e = est["c"][c["key"]]
        if not e["on"]:
            continue
        if c["tipo"] == "simple":
            out.append(c["text"] + (BOYA if e["buoy"] else ""))
        elif c["tipo"] == "group":
            for name in c["subs"]:
                if e["subs"][name]["on"]:
                    out.append(c["prefix"] + name + (BOYA if e["subs"][name]["buoy"] else ""))
        else:
            base = c["base"]
            if e["in"] and e["out"]:
                if e["in_buoy"] == e["out_buoy"]:
                    out.append(f"{base} IN/OUT" + (BOYA if e["in_buoy"] else ""))
                else:  # uno de los dos en boya: cada uno en su linea
                    out.append(f"{base} IN" + (BOYA if e["in_buoy"] else ""))
                    out.append(f"{base} OUT" + (BOYA if e["out_buoy"] else ""))
            elif e["in"]:
                out.append(f"{base} IN" + (BOYA if e["in_buoy"] else ""))
            elif e["out"]:
                out.append(f"{base} OUT" + (BOYA if e["out_buoy"] else ""))
    for x in est["extra"]:
        if x["on"]:
            out.append(x["t"] + (BOYA if x["buoy"] else ""))
    return out


def tiene(formato: str, estado: dict, marca: str) -> bool:
    """True si un concepto con esa marca (head_tally, shore_gangway) esta tildado."""
    est = normalizar(formato, estado)
    return any(c.get(marca) and est["c"][c["key"]]["on"] for c in CATALOGO[formato])


# ------------------------------------------------------------------ fechas --
_MESES = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def fecha_larga(d: date) -> tuple[str, str, str, str]:
    """('September', '25', 'th', '2026'): dia siempre con 2 digitos y su ordinal."""
    n = d.day
    suf = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return _MESES[d.month - 1], f"{n:02d}", suf, str(d.year)


# ---------------------------------------------------------------- documentos --
def _vineta(x: Xlsx, ref: str) -> None:
    """La vineta de la plantilla: texto enriquecido (un punto en una fuente y los espacios en otra) que ya esta en
    las cadenas compartidas del libro (la numero 4); se reusa tal cual para que se vea igual que el original."""
    c = x._cell(ref)
    x._wipe(c)
    c.set("t", "s")
    etree.SubElement(c, M + "v").text = "4"


def _excel(buque: str, agencia: str, lines: list[str]) -> bytes:
    x = Xlsx(TEMPLATES / "voucher.xlsx")
    x.set_text("A12", f"         BAHIA BLANCA FOR THE VESSEL {buque.upper()} - {agencia.upper()}")
    n = len(lines)
    for i in range(max(n, 17)):
        row = 14 + i
        if i < n:
            _vineta(x, f"A{row}")
            x.copy_style("A14", f"A{row}")
            x.set_text(f"B{row}", lines[i])
            x.copy_style("B14", f"B{row}")
            x.show_row(row, 15.9)
        else:
            x.set_text(f"A{row}", "")
            x.set_text(f"B{row}", "")
    ultima = max(35, 14 + n + 3)  # el area de impresion crece si hay mas conceptos que filas
    wbx = x.parts["xl/workbook.xml"].decode("utf-8")
    x.parts["xl/workbook.xml"] = re.sub(r"(SERVICIOS!\$A\$2:\$I\$)35", rf"\g<1>{ultima}", wbx).encode("utf-8")
    return x.to_bytes()


def _word(buque: str, agencia: str, lines: list[str], fecha: date) -> bytes:
    d = docx.Document(TEMPLATES / "voucher_oceanway.docx")
    mes, dia, suf, anio = fecha_larga(fecha)
    nombre = buque.upper()

    # encabezado: "MV <BUQUE> <tab> BAHIA BLANCA port, <mes> <dd><st/nd/rd/th>, <aaaa>"
    runs = d.paragraphs[0].runs
    assert runs[0].text == "MV " and runs[3].text == "BAHIA BLANCA" and runs[10].text == ", 202", "plantilla de voucher cambio"
    runs[1].text = nombre
    runs[5].text = f" {mes}"
    runs[7].text = dia
    runs[8].text, runs[9].text = suf[0], suf[1]
    runs[10].text, runs[11].text = f", {anio[:3]}", anio[3]

    # conceptos: un parrafo por linea, con el formato del primero de la plantilla
    cell = d.tables[0].rows[0].cells[0]
    paras = cell.paragraphs
    modelo = paras[1]._p
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    anterior = paras[0]._p
    for linea in lines:
        nuevo = copy.deepcopy(modelo)
        rs = nuevo.findall(qn("w:r"))
        for r in rs[1:]:
            nuevo.remove(r)
        for t in rs[0].findall(qn("w:t")):
            rs[0].remove(t)
        t = rs[0].makeelement(qn("w:t"), {})
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = "- " + linea
        rs[0].append(t)
        anterior.addnext(nuevo)
        anterior = nuevo

    # pie: "Master of M/V <BUQUE>"
    pie = d.paragraphs[4].runs
    assert pie[0].text.startswith("Master of M/V"), "plantilla de voucher cambio"
    pie[1].text = nombre
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def generar(formato: str, buque: str, agencia: str, estado: dict, fecha: date) -> tuple[bytes, str, str]:
    """(contenido, nombre de archivo, tipo MIME)."""
    lines = lineas(formato, estado)
    base = re.sub(r'[\\/:*?"<>|]+', " ", f"VOUCHER {buque} - {agencia}").strip()
    if formato == WORD:
        return (_word(buque, agencia, lines, fecha), base + ".docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    return (_excel(buque, agencia, lines), base + ".xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ------------------------------------------------------------ head tally --
ACTIVIDADES_TALLY = ["BERTHING", "LDG"]  # LDG = loading; se puede escribir otra
_ORDEN_TURNO = {"01-07": 0, "07-13": 1, "13-19": 2, "19-01": 3}


def horas_turno(turno: str) -> str:
    """'07-13' -> '0700-1300'; '19-01' -> '1900-0100'."""
    a, b = turno.split("-")
    return f"{a}00-{b}00"


def tally(buque: str, turnos: list[dict]) -> bytes:
    """Voucher del Head Tally (REG-AM-39): una fila por turno elegido, con la fecha solo en el primero de cada dia.
    `turnos` = [{'fecha': 'yyyy-mm-dd', 'turno': '07-13', 'act': 'LDG'}, ...]."""
    x = Xlsx(TEMPLATES / "voucher_tally.xlsx")
    x.set_text("D7", buque.upper())
    filas = sorted(turnos, key=lambda t: (t["fecha"], _ORDEN_TURNO.get(t["turno"], 9)))
    primera, ultima = 12, 38
    capacidad = ultima - primera  # la ultima fila de la plantilla es la del borde de abajo
    if len(filas) > capacidad:  # mas turnos que filas: se agregan filas con el formato de las del medio
        extra = len(filas) - capacidad
        x.insert_rows(ultima, extra)
        for r in range(ultima, ultima + extra):
            x.show_row(r, 18)
            for col in "BCDEF":
                x.copy_style(f"{col}37", f"{col}{r}")
    anterior = None
    for i, t in enumerate(filas):
        r = primera + i
        if t["fecha"] != anterior:
            x.set_date(f"B{r}", date.fromisoformat(t["fecha"]))
            x.copy_style("B12", f"B{r}")
            anterior = t["fecha"]
        if r != primera:  # la plantilla tiene alineaciones distintas segun la fila: se usa la de la segunda para todas
            for col in "CDEF":
                x.copy_style(f"{col}13", f"{col}{r}")
        x.set_number(f"C{r}", 1)
        x.set_text(f"D{r}", horas_turno(t["turno"]))
        x.set_text(f"E{r}", t.get("act") or "LDG")
    return x.to_bytes()
