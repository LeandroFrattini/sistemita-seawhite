"""Excel de SERENOS (recibos de pago a cada sereno), aparte de la proforma final.

Es el libro «SERENOS» de Administracion: una solapa de recibos por categoria (ACEITEROS = normal, CEREALEROS = insalubre y PELEGROSIDAD = peligroso),
RESULTADO (un renglon por recibo, con el CUIL), CUIL (la lista de serenos) y WATCHMAN (la misma grilla de la hoja SERENOS de la proforma final).

Cada recibo es de UN sereno, para UN tramo de 6 horas (00-06, 06-12, 12-18, 18-24): jornales por tipo de dia x precio, menos jubilacion (14 %),
contribucion solidaria (2 %) y obra social (3 %), mas el viatico por jornal. Todo con formulas, como el libro original, y con su valor guardado
para que se vea igual en Vista protegida. Los importes salen de lo cargado en el paso 9 (mismas tarifas, mismos dias, mismos turnos).
"""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from pathlib import Path

from lxml import etree
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

LOGO = Path(__file__).resolve().parent / "templates" / "logo_recibos.png"
VENTANAS = ("00-06", "06-12", "12-18", "18-24")
HOJAS = ("ACEITEROS", "CEREALEROS", "PELEGROSIDAD")
LUGAR = {"ACEITEROS": "INGENIERO  WHITE", "CEREALEROS": "BAHIA BLANCA", "PELEGROSIDAD": "INGENIERO  WHITE"}
DOMICILIO = "Facundo zuviria 401 - CUIT 30-70723233-9"
TEXTO_RECIBI = "Recibí de SEA WHITE S.A.  por cuenta y orden del Capitán por los jornal/es cumplidos a bordo del buque:"
RET_JUBIL, CONTRIB_SOL, OBRA_SOCIAL = 0.14, 2, 0.03   # los descuentos de cada recibo
MESES = ("ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE")

# ---------------------------------------------------------------------------------------------------- renglones de cada recibo
# (texto, tipos de dia que van en ese renglon, categoria). Tipos: LV (lunes a viernes), SAB (sabado), DOM (domingo, feriado y no laborable).
LINEAS = {
    "ACEITEROS": {
        "00-06": [("Inhabil s/insalubr", {"LV"}, "NORMAL"), ("Sabado S/insal", {"SAB"}, "NORMAL"), ("Dom/Feriado s/insal", {"DOM"}, "NORMAL")],
        "06-12": [("Habil s/insalubr", {"LV"}, "NORMAL"), ("Sabado Habil S/insal", {"SAB"}, "NORMAL"), ("Dom/Feriado s/insal", {"DOM"}, "NORMAL")],
        "12-18": [("Habil s/insalubr", {"LV"}, "NORMAL"), ("Sabado S/insal", {"SAB"}, "NORMAL"), ("Dom/Feriado s/insal", {"DOM"}, "NORMAL")],
        "18-24": [("Inhabil s/insalubr", {"LV"}, "NORMAL"), ("Sabado S/insal", {"SAB"}, "NORMAL"), ("Dom/Feriado s/insal", {"DOM"}, "NORMAL")]},
    "PELEGROSIDAD": {
        "00-06": [("Inhabil c/peligrosidad", {"LV"}, "PELIGROSO"), ("Sabado c/ peligrosidad", {"SAB"}, "PELIGROSO"), ("Dom/Feriado c/peligrosidad", {"DOM"}, "PELIGROSO")],
        "06-12": [("Habil c/peligrosidad", {"LV"}, "PELIGROSO"), ("Sabado c/ peligrosidad", {"SAB"}, "PELIGROSO"), ("Dom/Feriado c/peligrosidad", {"DOM"}, "PELIGROSO")],
        "12-18": [("Habil c/peligrosidad", {"LV"}, "PELIGROSO"), ("Sabado c/ peligrosidad", {"SAB"}, "PELIGROSO"), ("Dom/Feriado c/peligrosidad", {"DOM"}, "PELIGROSO")],
        "18-24": [("Inhabil c/peligrosidad", {"LV"}, "PELIGROSO"), ("Sabado c/ peligrosidad", {"SAB"}, "PELIGROSO"), ("Dom/Feriado c/peligrosidad", {"DOM"}, "PELIGROSO")]},
    "CEREALEROS": {
        "00-06": [("Inhabil/Sab s/insalubr", {"LV", "SAB"}, "NORMAL"), ("Inhabil/Sab C/insalubr", {"LV", "SAB"}, "INSALUBRE"),
                  ("Dom/Feriado s/insal", {"DOM"}, "NORMAL"), ("Dom/Feriado c/insal", {"DOM"}, "INSALUBRE")],
        "06-12": [("Habil/Sab s/insalubr", {"LV", "SAB"}, "NORMAL"), ("Habil/Sab C/insalubr", {"LV", "SAB"}, "INSALUBRE"),
                  ("Dom/Feriado s/insal", {"DOM"}, "NORMAL"), ("Dom/Feriado c/insal", {"DOM"}, "INSALUBRE")],
        "12-18": [("Habil s/insalubr", {"LV"}, "NORMAL"), ("Habil C/insalubr", {"LV"}, "INSALUBRE"), ("Sabado s/insalubr", {"SAB"}, "NORMAL"),
                  ("Sabado C/insalubr", {"SAB"}, "INSALUBRE"), ("Dom/Feriado s/insal", {"DOM"}, "NORMAL"), ("Dom/Feriado c/insal", {"DOM"}, "INSALUBRE")],
        "18-24": [("Inhabil/Sab s/insalubr", {"LV", "SAB"}, "NORMAL"), ("Inhabil/Sab C/insalubr", {"LV", "SAB"}, "INSALUBRE"),
                  ("Dom/Feriado s/insal", {"DOM"}, "NORMAL"), ("Dom/Feriado c/insal", {"DOM"}, "INSALUBRE")]},
}


def tier_de(clave: str, ventana: str) -> str:
    """Tarifa que corresponde a un tipo de dia en un tramo (mismas reglas que el paso 9): domingo/feriado; sabado solo es habil de 06 a 12; semana 06 a 18 es habil."""
    if clave == "DOM":
        return "DOMINGO"
    if clave == "SAB":
        return "HABIL" if ventana == "06-12" else "INHABIL"
    return "HABIL" if ventana in ("06-12", "12-18") else "INHABIL"


def _clave_dia(fila: dict) -> str:
    if fila["tier"] == "DOMINGO":
        return "DOM"
    return "SAB" if fila["tipo_dia"] == "SABADO" else "LV"


def hoja_de(cat: str, hay_insalubre: bool) -> str:
    """Solapa de recibos de un turno: peligroso -> PELEGROSIDAD, insalubre -> CEREALEROS y normal -> ACEITEROS
    (si el barco tiene turnos insalubres, los normales van en CEREALEROS, que trae los renglones «s/insalubr»)."""
    if cat == "PELIGROSO":
        return "PELEGROSIDAD"
    if cat == "INSALUBRE" or hay_insalubre:
        return "CEREALEROS"
    return "ACEITEROS"


def _r2(n: float) -> float:
    return round(n + 0.0, 2)


# ---------------------------------------------------------------------------------------------------------------- recibos
def faltantes(serenos: dict) -> list[str]:
    """Lo que impide armar los recibos: turnos sin sereno, sin categoria, o sin precio (devuelve frases para mostrar tal cual)."""
    out = []
    for f in serenos["filas"]:
        a, b = f["ventana"].split("-")
        donde = f"{f['fecha'][8:10]}/{f['fecha'][5:7]} de {a} a {b}"
        if not f.get("nombre"):
            out.append(f"Falta elegir el sereno del turno del {donde}.")
        if not f.get("cat_efectiva"):
            out.append(f"Falta la categoría del turno del {donde}.")
        elif not f["precio"]:
            out.append(f"El turno del {donde} no tiene precio.")
    if serenos["filas"] and "viatico" in serenos.get("faltan", []):
        out.append("Falta elegir el viático (White o Galván).")
    return out


def armar_recibos(serenos: dict, base: dict) -> list[dict]:
    """Un recibo por sereno, solapa, tramo de 6 horas y mes. Lanza ValueError si algo no cierra con el paso 9."""
    filas = serenos["filas"]
    hay_ins = any(f["cat_efectiva"] == "INSALUBRE" for f in filas)
    unit_viat = serenos["viatico"]["unit"] if serenos["viatico"]["tipo"] else 0.0
    grupos: dict[tuple, dict] = {}
    for f in filas:
        hoja = hoja_de(f["cat_efectiva"], hay_ins)
        clave = (hoja, f["ventana"], f["nombre"], f["fecha"][:7])
        r = grupos.get(clave)
        if r is None:
            r = grupos[clave] = {"hoja": hoja, "ventana": f["ventana"], "nombre": f["nombre"], "mes": date(int(f["fecha"][:4]), int(f["fecha"][5:7]), 1), "cuenta": {}, "extra": {}, "fechas": set()}
        r["fechas"].add(f["fecha"])
        k = _clave_dia(f)
        lineas = LINEAS[hoja][f["ventana"]]
        idx = next((i for i, (_, claves, cat) in enumerate(lineas) if k in claves and cat == f["cat_efectiva"]), None)
        if idx is None:
            raise ValueError(f"No hay renglón de recibo para un turno {f['cat_efectiva']} del {f['fecha']} ({f['ventana']}).")
        tarifa = _r2(base.get(("SERENO", f["cat_efectiva"], tier_de(k, f["ventana"])), 0.0))
        if abs(f["precio"] - tarifa) < 0.005:
            r["cuenta"][idx] = r["cuenta"].get(idx, 0) + 1
        else:  # el precio de ese turno se cambio a mano en el paso 9: va en un renglon aparte con su precio
            r["extra"][(idx, _r2(f["precio"]))] = r["extra"].get((idx, _r2(f["precio"])), 0) + 1
    recibos = []
    for r in grupos.values():
        lineas = []
        for i, (texto, claves, cat) in enumerate(LINEAS[r["hoja"]][r["ventana"]]):
            k0 = sorted(claves)[0] if len(claves) == 1 else "LV"
            tarifa = _r2(base.get(("SERENO", cat, tier_de(k0, r["ventana"])), 0.0))
            lineas.append({"texto": texto, "cant": r["cuenta"].get(i, 0), "precio": tarifa})
        for (i, precio), cant in sorted(r["extra"].items()):
            lineas.append({"texto": LINEAS[r["hoja"]][r["ventana"]][i][0] + " *", "cant": cant, "precio": precio})
        bruto = sum(l["cant"] * l["precio"] for l in lineas)
        jornales = sum(l["cant"] for l in lineas)
        ret, contrib, obra = round(bruto * RET_JUBIL, 2), bruto * CONTRIB_SOL / 100, round(bruto * OBRA_SOCIAL, 2)
        ded = ret + contrib + obra
        viat = jornales * unit_viat
        r.update(lineas=lineas, bruto=bruto, ret=ret, contrib=contrib, obra=obra, ded=ded, jornales=jornales, viat_unit=unit_viat, viat=viat, neto=bruto - ded + viat)
        recibos.append(r)
    orden_hoja = {h: i for i, h in enumerate(HOJAS)}
    recibos.sort(key=lambda r: (orden_hoja[r["hoja"]], r["mes"], VENTANAS.index(r["ventana"]), r["nombre"]))
    # control: lo que suman los recibos tiene que ser exactamente lo del paso 9
    if abs(sum(r["bruto"] for r in recibos) - serenos["subtotal"]) > 0.01 or sum(r["jornales"] for r in recibos) != serenos["cantidad"]:
        raise ValueError("Los recibos no coinciden con los importes del paso 9. No se generó el Excel para no entregar datos incorrectos.")
    return recibos


# ------------------------------------------------------------------------------------------------------------- escritura
SEGOE = Font(name="Segoe UI", size=12)
SEGOE_B = Font(name="Segoe UI", size=12, bold=True)
ARIAL12 = Font(name="Arial", size=12)
ARIAL12_B = Font(name="Arial", size=12, bold=True)
NUM2 = "#,##0.00"
FECHA_LARGA = '[$-2C0A]d" de "mmmm" de "yyyy;@'
FINA = Side(style="thin")
GRUESA = Side(style="medium")


class Hoja:
    """Hoja de openpyxl que ademas anota el valor de cada formula (se guarda en el archivo para Vista protegida)."""

    def __init__(self, ws, cache: dict):
        self.ws, self.cache = ws, cache

    def v(self, ref, valor, font=SEGOE, fmt=None, h=None, borde=None, wrap=False, v=None):
        c = self.ws[ref]
        c.value = valor
        c.font = font
        if fmt:
            c.number_format = fmt
        if h or v or wrap:
            c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
        if borde:
            c.border = borde
        return c

    def f(self, ref, formula, cached, font=SEGOE, fmt=None, h=None, borde=None):
        c = self.v(ref, "=" + formula, font, fmt, h, borde)
        self.cache[(self.ws.title, ref)] = cached
        return c


def _recibo(h: Hoja, S: int, r: dict, buque: str, hoy: date, viat_ref: str, n: int) -> int:
    """Escribe un recibo desde la fila S. Devuelve la fila del NETO. n es la cantidad de renglones de la solapa (para el alto)."""
    ws = h.ws
    ws.add_image(_logo(), f"B{S}")
    ws.row_dimensions[S + 5].height = 18
    h.v(f"B{S + 5}", DOMICILIO, SEGOE_B, h="left")
    h.v(f"F{S + 5}", LUGAR[r["hoja"]], SEGOE)
    h.v(f"I{S + 5}", hoy, SEGOE, FECHA_LARGA, h="left")
    ws.merge_cells(f"I{S + 5}:K{S + 5}")
    h.v(f"B{S + 7}", TEXTO_RECIBI, SEGOE, h="left")
    h.v(f"B{S + 8}", buque, SEGOE_B, h="left")
    h.v(f"B{S + 9}", "que corresponde a      ", SEGOE)
    h.v(f"C{S + 9}", r["nombre"], SEGOE_B)
    ws.merge_cells(f"C{S + 9}:F{S + 9}")
    a, b = r["ventana"].split("-")
    h.v(f"B{S + 10}", "Servicio comprendido en el horario de ", SEGOE)
    h.v(f"C{S + 10}", f"{a}00 a {b}00", SEGOE)
    L0 = S + 14
    for i, l in enumerate(r["lineas"]):
        fila = L0 + i
        h.v(f"C{fila}", l["texto"], SEGOE, h="left")
        if l["cant"]:
            h.v(f"F{fila}", l["cant"], SEGOE, "0", h="right")
        h.v(f"G{fila}", "JORNALES a", SEGOE, h="right")
        h.v(f"H{fila}", "$", SEGOE, h="right")
        h.v(f"I{fila}", l["precio"], SEGOE, NUM2)
        h.v(f"J{fila}", "$", SEGOE)
        h.f(f"K{fila}", f"F{fila}*I{fila}", l["cant"] * l["precio"], SEGOE, NUM2)
    ult = L0 + len(r["lineas"]) - 1
    ws[f"K{ult}"].border = Border(bottom=FINA)
    G0 = ult + 1
    h.f(f"K{G0}", f"SUM(K{L0 - 1}:K{ult})", r["bruto"], SEGOE, NUM2)
    filas_ded = [("Ret.Jubil/Ley 19032..........14%", f"ROUND(K{G0}*0.14,2)", r["ret"]),
                 ("Contribucion Solidaria....................2%", f"K{G0}*2/100", r["contrib"]),
                 ("Obra Social.........................3%", f"ROUND(K{G0}*0.03,2)", r["obra"])]
    for i, (texto, formula, valor) in enumerate(filas_ded):
        fila = G0 + 1 + i
        h.v(f"D{fila}", texto, SEGOE)
        h.v(f"H{fila}", "$", SEGOE, h="right")
        h.f(f"I{fila}", formula, valor, SEGOE, NUM2)
        h.v(f"J{fila}", "$", SEGOE)
    KD = G0 + 4
    h.f(f"K{KD}", f"SUM(I{G0}:I{KD - 1})", r["ded"], SEGOE, NUM2)
    fv = KD + 1
    h.v(f"D{fv}", "Viatico ", SEGOE)
    h.f(f"G{fv}", f"SUM(F{L0}:F{ult})", r["jornales"], SEGOE, "0", h="right")
    h.v(f"H{fv}", "$", SEGOE, h="right")
    h.f(f"I{fv}", viat_ref, r["viat_unit"], SEGOE, NUM2)
    h.f(f"K{fv}", f"+G{fv}*I{fv}", r["viat"], SEGOE, NUM2, borde=Border(bottom=FINA))
    fn = fv + 1
    h.v(f"J{fn}", "$", SEGOE)
    h.f(f"K{fn}", f"(K{G0}-K{KD}+K{fv})", r["neto"], SEGOE, NUM2)
    firma = S + 26 + n
    h.v(f"C{firma}", "Recibí copia del presente recibo", SEGOE, h="center", borde=Border(top=FINA))
    for col in "DEFG":
        ws[f"{col}{firma}"].border = Border(top=FINA)
    ws.merge_cells(f"C{firma}:G{firma}")
    if any(l["texto"].endswith(" *") for l in r["lineas"]):
        h.v(f"C{fn + 2}", "* Precio del turno modificado a mano en el paso 9 de la proforma final.", Font(name="Segoe UI", size=9, italic=True), h="left")
    r["celdas"] = {"nombre": f"C{S + 9}", "neto": f"K{fn}", "viat": f"K{fv}", "jornales": f"G{fv}", "bruto": f"K{G0}"}
    return fn


def _logo() -> XLImage:
    img = XLImage(str(LOGO))
    img.width, img.height = 232, 78
    return img


def _hoja_recibos(wb: Workbook, cache: dict, nombre: str, recibos: list[dict], buque: str, hoy: date, viat_unit: float) -> None:
    ws = wb.create_sheet(nombre)
    h = Hoja(ws, cache)
    for col, ancho in (("A", 2.55), ("B", 40.89), ("C", 14.11), ("D", 6.89), ("E", 8.43), ("F", 8.43), ("G", 15.89), ("H", 2.0), ("I", 15.78), ("J", 3.89), ("K", 15.55)):
        ws.column_dimensions[col].width = ancho
    n = max(len(r["lineas"]) for r in recibos)  # renglones del recibo mas largo (incluye los de «precio a mano»)
    paso = 29 + n
    ws.sheet_view.showGridLines = False
    h.v("J1", "Viático por jornal", SEGOE, h="right")
    h.v("K1", viat_unit, SEGOE, "0.00", h="right")
    for i, r in enumerate(recibos):
        S = 1 + i * paso
        for fila in range(S, S + paso):
            ws.row_dimensions[fila].height = 18
        _recibo(h, S, r, buque, hoy, "$K$1", n)
        if i % 2 == 1 and i < len(recibos) - 1:
            ws.row_breaks.append(Break(id=S + paso - 1))
    fin = max(1, len(recibos)) * paso
    ws.print_area = f"B1:K{fin}"
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.page_setup.scale = min(76, int(780 / (2 * paso * 18) * 100))
    ws.page_margins.left, ws.page_margins.right, ws.page_margins.top, ws.page_margins.bottom = 0.59, 0.2, 0.3, 0.3


def _resultado(wb: Workbook, cache: dict, recibos: list[dict], buque: str, padron: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("RESULTADO", 0)
    h = Hoja(ws, cache)
    for col, ancho in (("A", 22), ("B", 35.55), ("C", 12), ("D", 27.66), ("E", 18), ("F", 23.11), ("G", 16), ("H", 8)):
        ws.column_dimensions[col].width = ancho
    for col, t in zip("ABCDEFGH", ("MES", "BUQUE", "TURNO", "SERENOS", "CUIL", "TOTAL", "VIATICOS", "DIAS")):
        h.v(f"{col}1", t, ARIAL12, h="center", borde=Border(bottom=FINA))
    orden = sorted(recibos, key=lambda r: (r["mes"], VENTANAS.index(r["ventana"]), r["nombre"], HOJAS.index(r["hoja"])))
    cuil_de = {n: c for n, c in padron}
    for i, r in enumerate(orden):
        fila = 2 + i
        hoja, c = r["hoja"], r["celdas"]
        h.v(f"A{fila}", r["mes"], ARIAL12, "mmmm yyyy", h="center")
        h.v(f"B{fila}", buque, ARIAL12)
        a, b = r["ventana"].split("-")
        h.v(f"C{fila}", f"{a} - {b}", ARIAL12, h="center")
        h.f(f"D{fila}", f"{hoja}!{c['nombre']}", r["nombre"], ARIAL12)
        h.f(f"E{fila}", f'IFERROR(VLOOKUP(D{fila},CUIL!$B$5:$C$300,2,FALSE),"SIN CUIL")', cuil_de.get(r["nombre"], "SIN CUIL"), ARIAL12)
        h.f(f"F{fila}", f"{hoja}!{c['neto']}", r["neto"], ARIAL12, NUM2)
        h.f(f"G{fila}", f"{hoja}!{c['viat']}", r["viat"], ARIAL12, NUM2)
        h.v(f"H{fila}", 1, ARIAL12, "0")  # DIAS: siempre 1 por tramo (por facturacion), aunque haya trabajado ese tramo en dos dias; los jornales y los viaticos siguen sumando todos
    ult = 1 + len(orden)
    t = ult + 1
    h.v(f"D{t}", "TOTAL", ARIAL12_B, h="right")
    h.f(f"F{t}", f"SUM(F2:F{ult})", sum(r["neto"] for r in orden), ARIAL12_B, NUM2, borde=Border(top=FINA))
    h.f(f"G{t}", f"SUM(G2:G{ult})", sum(r["viat"] for r in orden), ARIAL12_B, NUM2, borde=Border(top=FINA))
    h.f(f"H{t}", f"SUM(H2:H{ult})", len(orden), ARIAL12_B, "0", borde=Border(top=FINA))
    # total de cada sereno (suma de todos sus recibos)
    k = t + 3
    h.v(f"D{k}", "TOTAL POR SERENO", ARIAL12_B)
    for col, tx in zip("DEFGH", ("SERENOS", "CUIL", "TOTAL", "VIATICOS", "DIAS")):
        h.v(f"{col}{k + 1}", tx, ARIAL12, h="center", borde=Border(bottom=FINA))
    nombres = sorted({r["nombre"] for r in orden})
    for j, nom in enumerate(nombres):
        fila = k + 2 + j
        mios = [r for r in orden if r["nombre"] == nom]
        h.v(f"D{fila}", nom, ARIAL12)
        h.f(f"E{fila}", f'IFERROR(VLOOKUP(D{fila},CUIL!$B$5:$C$300,2,FALSE),"SIN CUIL")', cuil_de.get(nom, "SIN CUIL"), ARIAL12)
        h.f(f"F{fila}", f"SUMIF($D$2:$D${ult},D{fila},F$2:F${ult})", sum(r["neto"] for r in mios), ARIAL12, NUM2)
        h.f(f"G{fila}", f"SUMIF($D$2:$D${ult},D{fila},G$2:G${ult})", sum(r["viat"] for r in mios), ARIAL12, NUM2)
        h.f(f"H{fila}", f"SUMIF($D$2:$D${ult},D{fila},H$2:H${ult})", len(mios), ARIAL12, "0")
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)  # que entren todas las columnas en una hoja al imprimir
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


def _cuil(wb: Workbook, cache: dict, padron: list[tuple[str, str]]) -> None:
    ws = wb.create_sheet("CUIL")
    h = Hoja(ws, cache)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 20
    h.v("B2", "SERENOS", Font(name="Calibri", size=16, bold=True, italic=True, underline="single"), h="center")
    ws.merge_cells("B2:C2")
    for col, t in (("B", "NOMBRE APELLIDO"), ("C", "CUIL")):
        c = h.v(f"{col}4", t, Font(name="Calibri", size=14, bold=True, color="FFFFFF"), h="center")
        c.fill = PatternFill("solid", fgColor="2F5597")
    for i, (n, c) in enumerate(sorted(padron)):
        h.v(f"B{5 + i}", n, Font(name="Calibri", size=13), borde=Border(bottom=FINA, left=FINA, right=FINA))
        h.v(f"C{5 + i}", c, Font(name="Calibri", size=13), h="center", borde=Border(bottom=FINA, left=FINA, right=FINA))


def _watchman(wb: Workbook, cache: dict, serenos: dict, buque: str, pct: float, dolar: float, sitio: str, feriados: set[str]) -> None:
    """La misma grilla de la hoja SERENOS de la proforma final: turnos x dias en pesos, subtotal, cargas sociales, viaticos y total (y en USD)."""
    ws = wb.create_sheet("WATCHMAN")
    h = Hoja(ws, cache)
    ws.sheet_view.showGridLines = False
    filas = serenos["filas"]
    fechas = sorted({f["fecha"] for f in filas})
    ncols = max(4, len(fechas))
    c0 = 4                                   # D: primer dia
    ctot = c0 + ncols                        # columna del TOTAL
    L = lambda n: ws.cell(row=1, column=n).column_letter
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 11
    for k in range(c0, ctot + 1):
        ws.column_dimensions[L(k)].width = 17
    caja = Border(left=GRUESA, right=GRUESA, top=GRUESA, bottom=GRUESA)
    cb = Font(name="Arial", size=11, bold=True)
    h.v("B2", buque, Font(name="Arial", size=13, bold=True))
    h.v("C4", "TURNO", Font(name="Arial", size=12, bold=True), h="center", borde=caja)
    h.v("B4", "SERENOS", Font(name="Arial", size=12, bold=True), h="center", borde=caja)
    for k in range(ncols):
        col = L(c0 + k)
        if k < len(fechas):
            d = date.fromisoformat(fechas[k])
            rojo = d.weekday() >= 5 or fechas[k] in feriados
            h.v(f"{col}4", d, Font(name="Arial", size=11, bold=True, color="C00000" if rojo else "000000"), "ddd dd/mm/yyyy", h="center", borde=caja)
        else:
            h.v(f"{col}4", None, cb, borde=caja)
    h.v(f"{L(ctot)}4", "TOTAL", cb, h="center", borde=caja)
    por_celda: dict[tuple[str, str], float] = {}
    nombres: dict[tuple[str, str], list[str]] = {}
    for f in filas:
        k = (f["fecha"], f["ventana"])
        por_celda[k] = por_celda.get(k, 0.0) + f["total"]
        if f.get("nombre"):
            nombres.setdefault(k, []).append(f["nombre"])
    r0 = 5
    dinero = '"$" #,##0.00'
    for b, v in enumerate(VENTANAS):
        ra = r0 + 3 * b
        quienes = []
        for fch in fechas:
            for n in nombres.get((fch, v), []):
                if n not in quienes:
                    quienes.append(n)
        h.v(f"B{ra}", " / ".join(quienes) if quienes else None, Font(name="Arial", size=10, bold=True), h="center", v="center", wrap=True, borde=caja)
        h.v(f"C{ra}", v, cb, h="center", v="center", borde=caja)
        for k in range(ncols):
            col = L(c0 + k)
            monto = por_celda.get((fechas[k], v)) if k < len(fechas) else None
            h.v(f"{col}{ra}", monto if monto else None, Font(name="Arial", size=11), dinero, h="center", v="center", borde=caja)
        tot = sum(por_celda.get((fch, v), 0.0) for fch in fechas)
        h.f(f"{L(ctot)}{ra}", f"SUM({L(c0)}{ra}:{L(ctot - 1)}{ra + 2})", tot, cb, dinero, h="center", borde=caja)
        ws[f"{L(ctot)}{ra}"].alignment = Alignment(horizontal="center", vertical="center")
        for col in ["B", "C"] + [L(c0 + k) for k in range(ncols)] + [L(ctot)]:
            ws.merge_cells(f"{col}{ra}:{col}{ra + 2}")  # (el borde de la celda de arriba se extiende a toda la union)
    sub = sum(por_celda.values())
    viat = serenos["viatico"]["total"] if serenos["viatico"]["tipo"] else 0.0
    lc, vc = L(ctot - 1), L(ctot)
    ini = r0 + 12

    def bloque(fila, etiqueta, valor, formula, fmt=dinero, fmt_et=None):
        h.v(f"{lc}{fila}", etiqueta, cb, fmt_et, h="center", v="center", borde=caja)
        if formula:
            h.f(f"{vc}{fila}", formula, valor, cb, fmt, h="center", borde=caja)
        else:
            h.v(f"{vc}{fila}", valor, cb, fmt, h="center", borde=caja)
        ws[f"{vc}{fila}"].alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"{lc}{fila}:{lc}{fila + 2}")
        ws.merge_cells(f"{vc}{fila}:{vc}{fila + 2}")

    bloque(ini, "SUBT", sub, f"SUM({vc}{r0}:{vc}{ini - 1})")
    bloque(ini + 3, pct / 100, sub * pct / 100, f"{vc}{ini}*{lc}{ini + 3}", fmt_et="0%")
    bloque(ini + 6, "CONV.", viat if viat else None, None)
    total = sub + sub * pct / 100 + viat
    bloque(ini + 9, "TOTAL", total, f"SUM({vc}{ini}:{vc}{ini + 8})")
    if dolar:
        bloque(ini + 12, "CAMBIO", dolar, None, fmt=NUM2)
        bloque(ini + 15, "TOTAL USD", total / dolar, f"{vc}{ini + 9}/{vc}{ini + 12}", fmt='"USD" #,##0.00')
    # quien cubrio cada turno, por dia
    k = ini
    h.v(f"B{k}", "SEA WHITE", Font(name="Arial", size=11, bold=True))
    k += 1
    if sitio:
        h.v(f"B{k}", sitio.upper(), Font(name="Arial", size=11))
        k += 1
    for fch in fechas:
        h.v(f"B{k}", date.fromisoformat(fch), Font(name="Arial", size=11, bold=True), "dddd dd/mm/yyyy")
        k += 1
        for f in filas:
            if f["fecha"] != fch:
                continue
            a, b2 = f["ventana"].split("-")
            h.v(f"B{k}", f"{a} A {b2}: {f['nombre']}", Font(name="Arial", size=11))
            k += 1
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.scale = 70


# ------------------------------------------------------------------------------------------------------------- armado
def _con_valores(datos: bytes, cache: dict, titulos: list[str]) -> bytes:
    """openpyxl guarda las formulas sin su resultado: se lo agrega a cada una (asi se ve bien en Vista protegida y en el celular)."""
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    entrada = zipfile.ZipFile(io.BytesIO(datos))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre in entrada.namelist():
            blob = entrada.read(nombre)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", nombre)
            if m:
                titulo = titulos[int(m.group(1)) - 1]
                raiz = etree.fromstring(blob)
                for c in raiz.iter(f"{{{ns}}}c"):
                    if c.find(f"{{{ns}}}f") is None:
                        continue
                    valor = cache.get((titulo, c.get("r")))
                    for v in c.findall(f"{{{ns}}}v"):
                        c.remove(v)
                    if valor is None:
                        continue
                    v = etree.SubElement(c, f"{{{ns}}}v")
                    if isinstance(valor, str):
                        c.set("t", "str")
                        v.text = valor
                    else:
                        c.attrib.pop("t", None)
                        v.text = format(float(valor), ".15g")
                blob = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
            z.writestr(nombre, blob)
    return salida.getvalue()


def generar(*, buque: str, sitio: str, serenos: dict, base: dict, padron: list[tuple[str, str]], pct: float, dolar: float, feriados: set[str], hoy: date) -> tuple[bytes, list[dict]]:
    """El libro SERENOS completo. `serenos` es lo que calcula personal.calcular_serenos (con el nombre de cada turno); `base`, las tarifas vigentes."""
    if not serenos["filas"]:
        raise ValueError("No hay turnos de serenos cargados en el paso 9.")
    problemas = faltantes(serenos)
    if problemas:
        raise ValueError(" ".join(problemas))
    recibos = armar_recibos(serenos, base)
    wb = Workbook()
    wb.remove(wb.active)
    cache: dict = {}
    unit = serenos["viatico"]["unit"] if serenos["viatico"]["tipo"] else 0.0
    por_hoja = {h: [r for r in recibos if r["hoja"] == h] for h in HOJAS}
    for nombre in HOJAS:
        if por_hoja[nombre]:  # una solapa sin recibos no se arma (como las hojas sin datos de la proforma final)
            _hoja_recibos(wb, cache, nombre, por_hoja[nombre], buque, hoy, unit)
    _resultado(wb, cache, recibos, buque, padron)
    _cuil(wb, cache, padron)
    _watchman(wb, cache, serenos, buque, pct, dolar, sitio, feriados)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return _con_valores(buf.getvalue(), cache, [ws.title for ws in wb.worksheets]), recibos
