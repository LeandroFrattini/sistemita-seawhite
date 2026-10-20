"""Excel con el resumen de un sereno: los barcos que atendio, sus jornales y lo que cobro en cada uno, con el bruto y el neto finales."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from .serenos_xlsx import _con_valores, MESES

NUM2 = "#,##0.00"
FINA = Side(style="thin")


def _mes(m: str) -> str:
    return f"{MESES[int(m[5:7]) - 1].capitalize()} {m[:4]}"


def _fecha(iso: str) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}" if iso else ""


def generar(*, nombre: str, cuil: str, mes: str, filas: list, hoy: date) -> bytes:
    """`filas`: renglones de SerenoPago (buque, mes, desde, hasta, jornales, bruto, descuentos, viaticos, neto)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    cache: dict = {}
    f12, f12b = Font(name="Arial", size=12), Font(name="Arial", size=12, bold=True)
    ws["A1"] = "RESUMEN DE SERENO"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"], ws["B2"] = "Sereno:", nombre
    ws["A3"], ws["B3"] = "CUIL:", cuil
    ws["A4"], ws["B4"] = "Período:", (_mes(mes) if mes else "Todos los meses")
    ws["A5"], ws["B5"] = "Emitido:", hoy.strftime("%d/%m/%Y")
    for r in range(2, 6):
        ws[f"A{r}"].font, ws[f"B{r}"].font = f12b, f12
    cab = ("BUQUE", "MES", "DESDE", "HASTA", "JORNALES", "BRUTO", "DESCUENTOS", "VIATICOS", "NETO")
    for i, t in enumerate(cab):
        c = ws.cell(row=7, column=1 + i, value=t)
        c.font, c.alignment, c.border = f12b, Alignment(horizontal="center"), Border(bottom=FINA)
    r = 8
    for f in filas:
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = f.buque, _mes(f.mes), _fecha(f.desde), _fecha(f.hasta)
        ws[f"E{r}"] = f.jornales
        ws[f"F{r}"], ws[f"G{r}"], ws[f"H{r}"], ws[f"I{r}"] = f.bruto, f.descuentos, f.viaticos, f.neto
        for col in "ABCDEFGHI":
            ws[f"{col}{r}"].font = f12
        for col in "FGHI":
            ws[f"{col}{r}"].number_format = NUM2
        ws[f"E{r}"].number_format = "0"
        r += 1
    ultimo = r - 1
    ws[f"A{r}"] = "TOTAL"
    ws[f"A{r}"].font = f12b
    for col, clave in (("E", "jornales"), ("F", "bruto"), ("G", "descuentos"), ("H", "viaticos"), ("I", "neto")):
        ws[f"{col}{r}"] = f"=SUM({col}8:{col}{ultimo})" if filas else 0
        c = ws[f"{col}{r}"]
        c.font, c.border = f12b, Border(top=FINA)
        c.number_format = "0" if col == "E" else NUM2
        if filas:
            cache[("RESUMEN", f"{col}{r}")] = sum(getattr(f, clave) for f in filas)
    ws[f"A{r + 2}"] = "BRUTO: lo que se gana por los jornales. NETO: lo que cobra (bruto menos jubilación, contribución solidaria y obra social, más los viáticos)."
    ws[f"A{r + 2}"].font = Font(name="Arial", size=10, italic=True)
    for col, ancho in zip("ABCDEFGHI", (30, 18, 13, 13, 11, 18, 16, 16, 18)):
        ws.column_dimensions[col].width = ancho
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    buf = io.BytesIO()
    wb.save(buf)
    return _con_valores(buf.getvalue(), cache, ["RESUMEN"])
