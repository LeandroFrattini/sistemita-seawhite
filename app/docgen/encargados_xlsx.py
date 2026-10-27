"""Excel de ENCARGADOS (Head Tally): la declaracion jurada del personal eventual contratado, con el formato del libro de Administracion.

Una hoja con el nombre del barco: encabezado (barco, operacion, tipo de carga, puerto, sitio, fecha) y, por dia, los turnos con su encargado:
jornal basico, insalubridad, 1era produccion, bruto y las retenciones (O.S. 3 %, jubilacion 11 %, sindical 3 %, ley 3 %) hasta el neto. Los aportes
de O.S., jubilacion y ley se calculan sobre el bruto hasta el tope del mes (el sindical se descuenta siempre). Al final: totales y viaticos.
Con formulas (como el original) y con su valor guardado para Vista protegida.
"""
from __future__ import annotations

import re
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.properties import PageSetupProperties

from ..encargados import PCT_JUBIL, PCT_LEY, PCT_OS, PCT_SIND
from .serenos_xlsx import _con_valores

PLANTILLA = Path(__file__).resolve().parent / "templates" / "encargados.xlsx"
HOJA_MODELO = "CL YINGNA HE"
COLS = "ABCDEFGHIJKLMNOP"
MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre")


def _horario(turno: str) -> str:
    a, b = turno.split("-")
    return f"{a}00-{b}00"


def _titulo_hoja(buque: str) -> str:
    t = re.sub(r"[\[\]\\/*?:]", " ", buque).strip()
    return (t or "ENCARGADOS")[:31]


def generar(*, buque: str, operacion: str, carga: str, sitio: str, calc: dict, cuil_de: dict[str, str], viatico_dia: float, hoy: date) -> bytes:
    filas = calc["filas"]
    wb = load_workbook(PLANTILLA)
    ws = wb[HOJA_MODELO]
    wb.remove(wb["Hoja1"])
    cache: dict = {}
    ws.title = _titulo_hoja(buque)
    # estilos de la plantilla, antes de borrar lo de ejemplo
    est = {k: {c: copy(ws[f"{c}{r}"]._style) for c in COLS + "Q"} for k, r in (("fecha", 20), ("turno", 21), ("blanca", 24), ("total", 31))}
    est_viat = {c: copy(ws[f"{c}33"]._style) for c in "MNOP"}
    est_blanco = copy(ws["R40"]._style)
    for rango in list(ws.merged_cells.ranges):
        if rango.min_row >= 17:
            ws.unmerge_cells(str(rango))
    for r in range(17, 90):
        for c in range(1, 19):
            cel = ws.cell(row=r, column=c)
            cel.value = None
            cel._style = copy(est_blanco)
    # encabezado
    ws["E5"] = buque
    ws["E7"] = operacion
    ws["E8"] = carga
    ws["E10"] = sitio
    ws["E11"] = hoy.strftime("%d/%m/%Y")
    ws["G16"] = calc.get("insal_pct", 0)
    # turnos por dia
    r = 20
    primero = r
    por_fecha: dict[str, list[dict]] = {}
    for f in filas:
        por_fecha.setdefault(f["fecha"], []).append(f)
    for fch, lista in por_fecha.items():
        for c in COLS + "Q":
            ws[f"{c}{r}"]._style = copy(est["fecha"][c])
        ws[f"B{r}"] = date.fromisoformat(fch)
        r += 1
        for f in lista:
            for c in COLS + "Q":
                ws[f"{c}{r}"]._style = copy(est["turno"][c])
            ws[f"A{r}"] = _horario(f["turno"])
            ws[f"B{r}"] = f["nombre"]
            ws[f"C{r}"] = int(re.sub(r"\D", "", cuil_de[f["nombre"]]))
            ws[f"D{r}"] = "ENCARGADO"
            for col, clave in (("F", "basico"), ("G", "insal"), ("H", "prod")):
                ws[f"{col}{r}"] = f[clave]
            ws[f"I{r}"] = 0
            ws[f"J{r}"] = 0
            k = f["bruto"]
            ws[f"K{r}"] = f"=SUM(F{r}:J{r})"
            cache[(ws.title, f"K{r}")] = f["basico"] + f["insal"] + f["prod"]
            k = cache[(ws.title, f"K{r}")]
            if f["estado"] == "completo":  # con el tope sin alcanzar: las formulas de siempre
                ws[f"L{r}"], ws[f"M{r}"], ws[f"O{r}"] = f"=K{r}*3%", f"=K{r}*11%", f"=K{r}*3%"
                cache[(ws.title, f"L{r}")], cache[(ws.title, f"M{r}")], cache[(ws.title, f"O{r}")] = k * PCT_OS, k * PCT_JUBIL, k * PCT_LEY
            else:                          # tope alcanzado (o a punto de alcanzarse): lo que corresponde descontar, sin la formula
                ws[f"L{r}"], ws[f"M{r}"], ws[f"O{r}"] = f["os"], f["jubil"], f["ley"]
            ws[f"N{r}"] = f"=K{r}*3/100"
            cache[(ws.title, f"N{r}")] = k * PCT_SIND
            ws[f"P{r}"] = f"=SUM(K{r}-L{r}-M{r}-N{r}-O{r})"
            cache[(ws.title, f"P{r}")] = f["neto"]
            f["_fila"] = r
            r += 1
        for c in COLS + "Q":
            ws[f"{c}{r}"]._style = copy(est["blanca"][c])
        r += 1
    ultimo = r - 1
    r += 1
    # totales
    for c in COLS + "Q":
        ws[f"{c}{r}"]._style = copy(est["total"][c])
    ws[f"D{r}"] = "TOTALES"
    sumas = {"F": "basico", "G": "insal", "H": "prod"}
    for col in "FGHIJKLMNOP":
        ws[f"{col}{r}"] = f"=SUM({col}{primero}:{col}{ultimo})"
        if col in sumas:
            v = sum(f[sumas[col]] for f in filas)
        elif col in "IJ":
            v = 0
        elif col == "K":
            v = sum(f["basico"] + f["insal"] + f["prod"] for f in filas)
        elif col == "L":
            v = sum(f["os"] for f in filas)
        elif col == "M":
            v = sum(f["jubil"] for f in filas)
        elif col == "N":
            v = sum(f["bruto"] * PCT_SIND for f in filas)
        elif col == "O":
            v = sum(f["ley"] for f in filas)
        else:
            v = sum(f["neto"] for f in filas)
        cache[(ws.title, f"{col}{r}")] = v
    # viaticos: uno por dia trabajado
    dias = len(por_fecha)
    r += 2
    for c in "MNOP":
        ws[f"{c}{r}"]._style = copy(est_viat[c])
    ws[f"M{r}"] = "Viaticos"
    ws.merge_cells(f"M{r}:O{r}")
    ws[f"P{r}"] = f"={viatico_dia:.2f}*{dias}"
    cache[(ws.title, f"P{r}")] = round(viatico_dia, 2) * dias
    fin = r
    # aviso de tope
    avisos = []
    for m, tope in calc["topes"].items():
        if any(f["fecha"][:7] == m and f["estado"] != "completo" for f in filas):
            avisos.append(f"Tope de sueldo bruto de {MESES[int(m[5:7]) - 1]} {m[:4]}: $ {tope:,.2f}. Alcanzado el tope del mes, O.S., Jubilación y Ley no se descuentan (el aporte sindical sí).".replace(",", "X").replace(".", ",").replace("X", "."))
    for nombre, m in calc.get("manuales", []):
        avisos.append(f"{nombre}: se marcó a mano que superó el tope de {MESES[int(m[5:7]) - 1]} {m[:4]}. O.S., Jubilación y Ley no se descuentan (el aporte sindical sí).")
    r += 2
    for a in avisos:
        ws[f"B{r}"] = a
        ws[f"B{r}"].font = Font(name="Century Gothic", size=12, italic=True)
        r += 1
    fin = max(fin, r - 1)
    ws.print_area = f"A1:P{fin}"
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.calculation.fullCalcOnLoad = True
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return _con_valores(buf.getvalue(), cache, [w.title for w in wb.worksheets])
