"""Un generador por documento. Cada uno abre su plantilla, completa solo las
celdas que antes salian por formula del libro viejo, y devuelve (nombre, bytes).
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import docx

from .fechas import en_ordinal_upper, en_parts_heinlein
from .xlsx_fill import Xlsx
from .xlsx_merge import merge_workbooks

TEMPLATES = Path(__file__).parent / "templates"

# clave, nombre largo, codigo REG-AM
DOCS: list[tuple[str, str, str]] = [
    ("dec_migra", "Declaración de Migraciones", "REG-AM-11"),
    ("dec_ana", "Declaración General Aduana (DEC. ANA)", "REG-AM-19"),
    ("permanencia", "Permanencia y Navegación", "REG-AM-20"),
    ("calados_ent", "Calados Entrada", "REG-AM-21"),
    ("rancho", "Recepción de Lista de Rancho", "REG-AM-23"),
    ("serenos", "Pedido de Serenos", "REG-AM-09"),
    ("ped_carga", "Pedido de Carga", "REG-AM-25"),
    ("bill", "Autorización BILL OF LADING", "REG-AM-26"),
    ("decla_pna", "Declaración General (DECLA PNA)", "REG-AM-10"),
    ("boyado", "Declaración Jurada de Boyas (BOYADO)", "REG-AM-18"),
    ("malvinas", "Declaración Jurada Malvinas", "Anexo 2"),
    ("pbip_entrada", "PBIP Entrada", "REG-AM-14"),
    ("pbip_salida", "PBIP Salida", "REG-AM-15"),
    ("free_damage", "Free Damage (Safety Stowage)", "Certificado de estiba"),
    ("seaworthy", "Seaworthy Certificate", "Certificado"),
]
DOC_KEYS = [d[0] for d in DOCS]

# solapa de la pantalla en la que sale cada documento
STAGES = {k: "ENTRADA" for k in DOC_KEYS}
STAGES["free_damage"] = "SALIDA"
STAGES["seaworthy"] = "SALIDA"
# documentos que no salen tildados por defecto (los que a veces piden y a veces no)
NOT_DEFAULT = set(DOC_KEYS)  # ahora ninguno sale tildado: se tildan solo los que se van a sacar
# los que a veces piden y a veces no (van bajo "A pedido")
OPTIONAL = {"decla_pna", "boyado", "malvinas", "pbip_entrada", "pbip_salida"}
# documentos que no tienen fecha para completar
NO_DATE = {"serenos", "pbip_entrada", "pbip_salida"}

SHORT_NAMES = {
    "dec_migra": "DEC. MIGRA", "dec_ana": "DEC. ANA", "permanencia": "Permanencia",
    "calados_ent": "Calados Entrada", "rancho": "Rancho", "serenos": "Serenos",
    "ped_carga": "Pedido de Carga", "bill": "BILL", "decla_pna": "DECLA PNA",
    "boyado": "BOYADO", "malvinas": "MALVINAS", "pbip_entrada": "PBIP ENTRADA", "pbip_salida": "PBIP SALIDA",
    "free_damage": "FREE DAMAGE", "seaworthy": "SEAWORTHY",
}

KIND_LETTER = {"BULK_CARRIER": "V", "TANKER": "T"}  # MV / MT


@dataclass
class Ctx:
    vessel: dict
    call: dict
    agency: object | None  # DocAgency o None
    dates: dict[str, date] = field(default_factory=dict)
    serenos: dict[str, date | None] = field(default_factory=dict)
    migra_modo: str = "ENTRADA"  # ENTRADA | SALIDA
    pna_modo: str = "ENTRADA"  # ENTRADA | SALIDA (selector de DECLA PNA)
    certificados: dict = field(default_factory=dict)  # vencimientos guardados con el buque
    firma: bytes | None = None  # firma del capitan de esta escala (PNG ya preparado)
    firma_keys: set = field(default_factory=set)  # documentos en los que se pega

    @property
    def name(self) -> str:
        return (self.vessel.get("name") or "").strip()

    @property
    def letter(self) -> str:
        return KIND_LETTER.get(self.vessel.get("kind"), "V")

    def agency_word(self, key: str) -> bool:
        """True si la agencia elegida usa su propio formato en Word para este documento."""
        a = self.agency
        if not a:
            return False
        if key in ("ped_carga", "bill") and a.doc_format == "HEINLEIN":
            return True
        return key in [w for w in (getattr(a, "word_docs", "") or "").split(",") if w]

    @property
    def is_heinlein(self) -> bool:
        return bool(self.agency) and self.agency.doc_format == "HEINLEIN"

    def v(self, key: str) -> str:
        return (self.vessel.get(key) or "").strip()

    def c(self, key: str) -> str:
        return (self.call.get(key) or "").strip()


def _agency_lines(agency) -> list[str]:
    return [l.strip() for l in (agency.address or "").splitlines() if l.strip()]


# ------------------------------------------------------------- Excel -------
def gen_dec_ana(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "dec_ana.xlsx")
    x.set_date("M4", ctx.dates.get("dec_ana"))
    x.set_text("K7", ctx.name)
    x.set_text("C8", ctx.v("flag"))
    x.set_text("K8", ctx.c("capitan"))
    x.set_text("N9", ctx.c("ultimo_puerto"))
    x.set_text("N37", ctx.c("terminal"))
    x.set_value("M39", ctx.v("trn"))
    return x.to_bytes()


def gen_dec_migra(ctx: Ctx) -> bytes:
    entrada = ctx.migra_modo != "SALIDA"
    x = Xlsx(TEMPLATES / "dec_migra.xlsx")
    d = ctx.dates.get("dec_migra")
    x.set_text("D4", "X Entrada" if entrada else "X Salida")
    # resultados guardados de las formulas que dependen del selector D4
    x.set_formula_result("D5", "2.Puerto de Entrada" if entrada else "2.Puerto de Salida")
    x.set_formula_result("E5", "3. Fecha y Hora de Entrada" if entrada else "3. Fecha y Hora de Salida")
    x.set_formula_result("E7", "6.Puerto de procedencia" if entrada else "6.Puerto de destino")
    x.set_formula_result("D13", "11. Calados de Entrada" if entrada else "11. Calados de Salida")
    x.set_text("AB7", ctx.c("procedencia"))
    x.set_text("AB8", ctx.c("destino"))
    x.set_formula_result("E8", ctx.c("procedencia") if entrada else ctx.c("destino"))

    x.set_text("A6", ctx.name)
    x.set_text("D6", "Bahia Blanca")
    x.set_date("E6", d)
    x.set_text("A8", ctx.v("flag"))
    x.set_text("D8", ctx.c("capitan"))
    x.set_value("A10", ctx.v("matricula"))
    x.set_text("C10", ctx.v("puerto_registro"))
    x.set_value("A12", ctx.v("trb"))
    x.set_value("C12", ctx.v("trn"))
    x.set_text("A14", ctx.c("terminal"))
    x.set_text("A16", ctx.c("descripcion_viaje"))
    x.set_text("A19", ctx.c("carga_detalle"))
    x.set_value("A22", ctx.c("tripulantes"))
    x.set_value("B22", ctx.c("pasajeros"))
    x.set_text("B27", ctx.c("lista_pasajeros"))
    x.set_date("D32", d)
    x.set_text("B52", ctx.c("destino"))
    return x.to_bytes()


def gen_permanencia(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "permanencia.xlsx")
    for ref in ("G12", "G13"):
        x.set_text(ref, ctx.name)
    for ref in ("L12", "L13"):
        x.set_text(ref, ctx.v("flag"))
    for ref in ("B15", "B16"):
        x.set_text(ref, ctx.c("terminal"))
    for ref in ("B18", "C19"):
        x.set_value(ref, ctx.c("estadia"))
    x.set_date("A23", ctx.dates.get("permanencia"))
    return x.to_bytes()


def gen_calados(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "calados_ent.xlsx")
    d = ctx.dates.get("calados_ent")
    x.set_text("B9", ctx.name)
    x.set_text("E9", ctx.v("flag"))
    x.set_date("E32", d)
    x.set_date("D45", d)
    return x.to_bytes()


def gen_rancho(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "rancho.xlsx")
    x.set_date("G8", ctx.dates.get("rancho"))
    return x.to_bytes()


# cada bloque del pedido de serenos: fila de la fecha de solicitud, de buque/bandera,
# de muelle y de la fecha de inicio del turno
_SERENO_BLOCKS = [
    ("t0006", 7, 11, 13, 15),
    ("t0612", 58, 62, 64, 66),
    ("t1218", 108, 112, 114, 116),
    ("t1824", 158, 162, 164, 166),
]


def gen_serenos(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "serenos.xlsx")
    for turno, r_sol, r_buque, r_muelle, r_turno in _SERENO_BLOCKS:
        x.set_date(f"F{r_sol}", ctx.serenos.get("solicitud"))
        x.set_text(f"B{r_buque}", ctx.name)
        x.set_text(f"G{r_buque}", ctx.v("flag"))
        x.set_text(f"B{r_muelle}", ctx.c("terminal"))
        x.set_date(f"D{r_turno}", ctx.serenos.get(turno))
    return x.to_bytes()


def _agency_logo(x: Xlsx, ctx: Ctx, box, box_w, box_h, pad: float = 10.0) -> None:
    a = ctx.agency
    if a is not None and a.logo:
        x.add_logo(a.logo, a.logo_mime, box, box_w, box_h, pad_pt=pad)


def gen_ped_carga(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "ped_carga.xlsx")
    x.set_text("K7", en_ordinal_upper(ctx.dates["ped_carga"]) if ctx.dates.get("ped_carga") else "")
    x.set_text("A10", ctx.c("exportador"))
    x.set_text("A11", ctx.c("ciudad_exportador"))
    x.set_text("H18", ctx.c("carga"))
    x.set_text("I29", ctx.name)
    _agency_logo(x, ctx, (1, 1, 4, 5), 124.2, 67.2)
    return x.to_bytes()


def gen_bill(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "bill.xlsx")
    x.set_text("E7", en_ordinal_upper(ctx.dates["bill"]) if ctx.dates.get("bill") else "")
    lines = _agency_lines(ctx.agency) if ctx.agency else []
    x.set_text("A10", ctx.agency.name if ctx.agency else "")
    x.set_text("A11", lines[0] if lines else "")
    x.set_text("A12", ", ".join(lines[1:]) if len(lines) > 1 else "")
    x.set_text("A22", ctx.v("flag"))
    x.set_text("D22", ctx.name)
    x.set_text("E31", ctx.name)
    x.set_text("D32", ctx.c("capitan"))
    _agency_logo(x, ctx, (1, 1, 2, 5), 145.2, 72.6)
    return x.to_bytes()


def _cert(x: Xlsx, ref: str, raw: str) -> None:
    """Vencimiento de certificado: fecha si viene como AAAA-MM-DD, si no tal cual."""
    raw = (raw or "").strip()
    try:
        x.set_date(ref, date.fromisoformat(raw)) if raw else x.set_text(ref, "")
    except ValueError:
        x.set_text(ref, raw)


# celda de DECLA PNA de cada certificado
_PNA_CERTS = [
    ("iapp", "G36"), ("radio", "G38"), ("equipo", "G40"), ("francobordo", "G42"),
    ("construccion", "G44"), ("desratizacion", "G46"), ("polucion", "G48"), ("cgs", "G50"),
    ("doc", "G52"), ("isps", "G54"), ("mlc", "G60"), ("fitness", "G62"),
    ("sewage", "G64"), ("clc", "G66"),
]


def gen_decla_pna(ctx: Ctx) -> bytes:
    entrada = ctx.pna_modo != "SALIDA"
    x = Xlsx(TEMPLATES / "decla_pna.xlsx")
    d = ctx.dates.get("decla_pna")
    x.set_text("D5", "X Entrada" if entrada else "X Salida")
    # resultados guardados de las formulas que dependen del selector D5
    x.set_formula_result("D6", "2.Puerto de Entrada" if entrada else "2.Puerto de Salida")
    x.set_formula_result("F6", "3. Fecha y Hora de Entrada" if entrada else "3. Fecha y Hora de Salida")
    x.set_formula_result("D8", "Puerto de procedencia" if entrada else "Puerto de destino")
    x.set_text("AA9", ctx.c("procedencia"))
    x.set_text("AA10", ctx.c("destino"))
    x.set_formula_result("D9", ctx.c("procedencia") if entrada else ctx.c("destino"))

    x.set_text("A7", ctx.name)
    x.set_date("F7", d)
    x.set_text("A9", ctx.v("flag"))
    x.set_text("B9", ctx.c("capitan"))
    x.set_text("A11", ctx.vessel.get("imo") or "")  # numero del certificado de matricula = IMO
    x.set_text("C11", ctx.v("puerto_registro"))
    x.set_value("A13", ctx.v("trb"))
    x.set_value("C13", ctx.v("trn"))
    x.set_text("A15", ctx.c("terminal"))
    x.set_text("A18", ctx.c("descripcion_viaje"))
    x.set_text("A21", ctx.c("carga_detalle"))
    x.set_value("A25", ctx.c("tripulantes"))
    x.set_value("B25", ctx.c("pasajeros"))
    x.set_text("B31", ctx.c("lista_pasajeros"))
    x.set_date("D32", d)

    # para uso oficial
    x.set_text("C38", ctx.v("call_sign"))
    x.set_value("E38", ctx.v("eslora"))
    x.set_text("C40", ctx.v("clasificacion"))
    x.set_value("E40", ctx.v("manga"))
    x.set_value("C42", ctx.v("velocidad"))
    x.set_value("E42", ctx.v("puntal"))
    x.set_text("C46", ctx.c("calado_proa"))
    x.set_text("C48", ctx.c("calado_popa"))
    x.set_text("C50", ctx.c("practico"))
    x.set_text("C52", ctx.c("remolque_proa"))
    x.set_text("C54", ctx.c("remolque_popa"))
    x.set_text("C56", ctx.c("estima"))
    x.set_text("C58", ctx.c("ultimo_puerto"))
    x.set_text("D60", ctx.c("calado_max"))  # como texto: la celda original redondea a entero
    x.set_text("C67", ctx.v("armador"))

    certs = dict(ctx.certificados or {})
    certs["iapp"] = certs.get("polucion", "")  # igual que la planilla original: IAPP = POLUCION
    for ref in ("G62", "G64", "G66"):  # estas celdas no traian formato de fecha
        x.copy_style("G60", ref)
    for key, ref in _PNA_CERTS:
        _cert(x, ref, certs.get(key, ""))
    x.set_text("G56", ctx.vessel.get("imo") or "")  # IMO e INMARSAT salen de los datos del buque
    x.set_text("G58", ctx.v("inmarsat"))
    return x.to_bytes()


def _matricula(ctx: Ctx) -> str:
    """El numero de matricula que piden las planillas de Prefectura es el IMO del buque."""
    return ctx.vessel.get("imo") or ""


def gen_boyado(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "boyado.xlsx")
    x.set_date("L7", ctx.dates.get("boyado"))
    x.set_text("D11", ctx.c("capitan"))
    x.set_text("O11", ctx.name)
    x.set_text("B14", ctx.v("flag"))
    x.set_value("F14", _matricula(ctx))
    x.set_text("D26", ctx.c("terminal"))
    return x.to_bytes()


def gen_malvinas(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "malvinas.xlsx")
    imo = _matricula(ctx)
    owner = ctx.v("armador")
    company = ctx.v("company_id")
    d = ctx.dates.get("malvinas")
    x.set_text("C20", ctx.c("capitan"))
    x.set_text("D24", ctx.name)
    x.set_value("H24", imo)
    x.set_value("B28", imo)
    x.set_text("G28", owner)
    x.set_value("D32", company)
    # las celdas en ingles repiten a las de arriba con una formula: se deja y se guarda su resultado
    for ref, val in (("C22", ctx.c("capitan")), ("D26", ctx.name), ("H26", imo),
                     ("B30", imo), ("G30", owner), ("H32", company)):
        x.set_formula_result(ref, val)
    x.set_text("B55", "bahia blanca" + (f" {d:%d-%m-%y}" if d else ""))
    return x.to_bytes()


def _gen_pbip(template: str):
    def gen(ctx: Ctx) -> bytes:
        x = Xlsx(TEMPLATES / template)
        x.set_text("F5", ctx.name)
        x.set_text("F6", ctx.v("flag"))
        x.set_text("F7", ctx.v("call_sign"))
        x.set_value("F8", _matricula(ctx))
        x.set_value("F9", _matricula(ctx))
        return x.to_bytes()
    return gen


gen_pbip_entrada = _gen_pbip("pbip_entrada.xlsx")
gen_pbip_salida = _gen_pbip("pbip_salida.xlsx")


# ---------------------------------------------------------- salida ---------
STOWAGE_TEXT = ("I hereby confirm that the loading operations of {carga} in bulk, carried out at this port, "
                "has been done under my ")


def _salida_header(x: Xlsx, ctx: Ctx, key: str) -> None:
    x.set_date("G7", ctx.dates.get(key))
    x.set_text("E21", ctx.name)
    _agency_logo(x, ctx, (1, 1, 2, 4), 111.6, 54.0, pad=8.0)


def gen_free_damage(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "free_damage.xlsx")
    _salida_header(x, ctx, "free_damage")
    x.set_text("A9", STOWAGE_TEXT.format(carga=ctx.c("carga") or "__________"))
    return x.to_bytes()


def gen_seaworthy(ctx: Ctx) -> bytes:
    x = Xlsx(TEMPLATES / "seaworthy.xlsx")
    _salida_header(x, ctx, "seaworthy")
    return x.to_bytes()


def gen_free_damage_word(ctx: Ctx) -> bytes:
    """Free Damage en el formato propio de la agencia (Word). El texto y los datos del buque
    se cambian dentro de las mismas partes del texto, asi queda todo con el formato original."""
    doc = docx.Document(TEMPLATES / "free_damage_word.docx")
    month, day, suffix, year = _heinlein_date(ctx.dates.get("free_damage"), zero_pad=False)
    blank = not year
    mv = f"M{ctx.letter}"

    r = _runs(doc, 0, "PUKA")  # MV PUKA   BAHIA BLANCA port, Septemeber 19th, 2026
    r[0].text, r[1].text = f"{mv} ", ctx.name
    r[5].text, r[7].text, r[8].text, r[9].text = month, day, ("" if blank else f"{suffix}, {year}"), ""
    if blank:
        r[6].text = ""

    r = _runs(doc, 3, "SINGAPUR")  # Captain of the SINGAPUR flag MV TEXEL ISLAND ... on Septemeber 19th, 2026
    r[1].text = ctx.v("flag")
    r[3].text, r[4].text = f"{mv} ", ctx.name
    r[9].text, r[11].text = month, day
    r[12].text, r[13].text, r[14].text = ("" if blank else f"{suffix},"), ("" if blank else f" {year}"), ""
    if blank:
        r[10].text = ""

    r = _runs(doc, 4, "Septemeber")  # I sign this free of damage certificate ... on Septemeber 19th, 2026.
    r[3].text, r[4].text, r[5].text = month, ("" if blank else f" {day}{suffix},"), ("" if blank else f" {year}")

    r = _runs(doc, 13, "TEXEL")  # Master of M/V TEXEL ISLAND
    r[0].text, r[2].text = f"Master of M/{ctx.letter}", ctx.name
    return _save(doc)


# ------------------------------------------------------------- Word --------
def _runs(doc, idx: int, expect: str):
    p = doc.paragraphs[idx]
    if expect not in p.text:
        raise RuntimeError(f"La plantilla de Heinlein cambió (párrafo {idx}: se esperaba «{expect}»)")
    return p.runs


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


BLANK_DATE_LINE = "_" * 22  # fecha en blanco en los Word: queda una linea para escribirla a mano


def _heinlein_date(d: date | None, zero_pad: bool = True) -> tuple[str, str, str, str]:
    if d is None:
        return BLANK_DATE_LINE, "", "", ""
    month, day, suffix, year = en_parts_heinlein(d)
    return month, (day if zero_pad else str(int(day))), suffix, year


def gen_heinlein_bl(ctx: Ctx) -> bytes:
    doc = docx.Document(TEMPLATES / "heinlein_bl.docx")
    month, day, suffix, year = _heinlein_date(ctx.dates.get("bill"))

    r = _runs(doc, 4, "TAMBLING")           # M.T. “TAMBLING”
    r[1].text = ctx.letter
    r[3].text = ctx.name

    r = _runs(doc, 6, "Bahia Blanca")       # Port of Bahia Blanca, October 04th, 2026.-
    r[4].text, r[6].text, r[7].text = month, day, suffix
    r[8].text, r[9].text, r[10].text = (f", {year}" if year else ""), "", ""

    a = ctx.agency
    lines = _agency_lines(a) if a else []
    _runs(doc, 9, "MARITIMA")[0].text = a.name if a else ""
    r = _runs(doc, 10, "PERU")
    r[0].text, r[1].text = (lines[0] if lines else ""), ""
    _runs(doc, 11, "Buenos Aires")[0].text = lines[1] if len(lines) > 1 else ""

    r = _runs(doc, 19, "TAMBLING")          # M/T “TAMBLING”
    r[1].text = ctx.letter
    r[3].text = ctx.name
    return _save(doc)


def gen_heinlein_cargo(ctx: Ctx) -> bytes:
    doc = docx.Document(TEMPLATES / "heinlein_cargo.docx")
    month, day, suffix, year = _heinlein_date(ctx.dates.get("ped_carga"))

    r = _runs(doc, 3, "TAMBLING")           # M.V. “TAMBLING”
    r[0].text = f"M.{ctx.letter}. “"
    r[1].text = ctx.name

    r = _runs(doc, 5, "Bahia Blanca")       # Port of Bahia Blanca, October 04th., 2026.-
    r[4].text, r[6].text = month, f"{day}{suffix}"
    r[7].text, r[8].text, r[9].text = (f"., {year}" if year else ""), "", ""

    r = _runs(doc, 8, "ACEITERA")           # exportador
    r[0].text, r[1].text = ctx.c("exportador"), ""

    r = _runs(doc, 13, "TAMBLING")          # the Singapore flag M.T. “TAMBLING” ...
    r[1].text = ctx.v("flag")
    r[3].text = ctx.letter
    r[5].text = ctx.name

    r = _runs(doc, 21, "TAMBLING")          # Master of M/T “TAMBLING”
    r[2].text = f"{ctx.letter} “"
    r[3].text = ctx.name

    cargo_cell = doc.tables[0].rows[1].cells[4]
    cargo_cell.paragraphs[0].runs[0].text = "    " + ctx.c("carga")
    return _save(doc)


# ---------------------------------------------------------- entrada --------
_XLSX_GEN = {
    "dec_migra": gen_dec_migra, "dec_ana": gen_dec_ana, "permanencia": gen_permanencia,
    "calados_ent": gen_calados, "rancho": gen_rancho, "serenos": gen_serenos,
    "ped_carga": gen_ped_carga, "bill": gen_bill, "decla_pna": gen_decla_pna,
    "boyado": gen_boyado, "malvinas": gen_malvinas,
    "pbip_entrada": gen_pbip_entrada, "pbip_salida": gen_pbip_salida,
    "free_damage": gen_free_damage, "seaworthy": gen_seaworthy,
}


def _safe(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "-", text).strip() or "buque"


# Donde va la firma del capitan en cada planilla: (caja de celdas, ancho y alto de la caja en puntos
# medidos en Excel, alineacion, alto maximo de la firma). Los Word de Heinlein no llevan firma.
FIRMA_BOXES: dict[str, tuple[tuple[int, int, int, int], float, float, str, float]] = {
    "dec_migra": ((5, 31, 5, 32), 194.4, 28.2, "right", 26),
    "dec_ana": ((2, 53, 5, 55), 147.0, 45.0, "center", 40),
    "permanencia": ((10, 19, 13, 21), 154.2, 57.6, "center", 46),
    "calados_ent": ((2, 33, 2, 35), 144.0, 45.0, "center", 40),
    "rancho": ((4, 37, 5, 40), 123.6, 60.0, "center", 46),
    "decla_pna": ((7, 32, 7, 33), 108.6, 38.4, "center", 34),
    "boyado": ((6, 29, 12, 31), 224.4, 39.6, "center", 36),
    "malvinas": ((7, 52, 8, 55), 212.4, 52.8, "center", 46),
    "pbip_entrada": ((3, 60, 8, 62), 372.0, 45.0, "center", 40),
    "pbip_salida": ((3, 41, 8, 43), 429.6, 43.2, "center", 40),
    "free_damage": ((4, 16, 8, 20), 274.2, 78.0, "center", 56),
    "seaworthy": ((4, 16, 8, 20), 274.2, 78.0, "center", 56),
}
FIRMA_DOCS = set(FIRMA_BOXES)


def _firmar(data: bytes, key: str, ctx: Ctx) -> bytes:
    box, w, h, align, max_h = FIRMA_BOXES[key]
    x = Xlsx(io.BytesIO(data))
    x.add_logo(ctx.firma, "image/png", box, w, h, pad_pt=2, tag="firma_capitan", name="Firma del capitan",
               shape_id=9002, align_x=align, max_h_pt=max_h)
    return x.to_bytes()


def build(key: str, ctx: Ctx) -> tuple[str, bytes]:
    base = SHORT_NAMES[key]
    suffix = _safe(ctx.name)
    if key == "free_damage" and ctx.agency_word("free_damage"):
        return f"Free Damage (formato agencia) - {suffix}.docx", gen_free_damage_word(ctx)
    if key in ("ped_carga", "bill") and ctx.is_heinlein:
        if key == "ped_carga":
            return f"Cargo Declaration (Heinlein) - {suffix}.docx", gen_heinlein_cargo(ctx)
        return f"Authorization BL (Heinlein) - {suffix}.docx", gen_heinlein_bl(ctx)
    data = _XLSX_GEN[key](ctx)
    if ctx.firma and key in ctx.firma_keys and key in FIRMA_BOXES:
        data = _firmar(data, key, ctx)
    return f"{base} - {suffix}.xlsx", data


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def build_many(keys: list[str], ctx: Ctx, single_workbook: bool = False) -> tuple[str, bytes, str]:
    """Devuelve (nombre, contenido, mime). Uno solo sale suelto; varios, en un zip.

    Con single_workbook, los Excel salen juntos en un solo libro (una hoja por
    documento) para imprimir todo de una vez; los Word de Heinlein no se pueden
    unir y van aparte, en un zip junto con ese libro.
    """
    files = [(k, *build(k, ctx)) for k in keys]
    if single_workbook:
        sheets = [(SHORT_NAMES[k], data) for k, name, data in files if name.endswith(".xlsx")]
        others = [(name, data) for k, name, data in files if not name.endswith(".xlsx")]
        if len(sheets) >= 2:
            book = (f"Documentacion - {_safe(ctx.name)}.xlsx", merge_workbooks(sheets))
            if not others:
                return book[0], book[1], XLSX_MIME
            files = [(None, *book)] + [(None, n, d) for n, d in others]
            return _zip(ctx, files)
    if len(files) == 1:
        _, name, data = files[0]
        return name, data, DOCX_MIME if name.endswith(".docx") else XLSX_MIME
    return _zip(ctx, files)


def _zip(ctx: Ctx, files) -> tuple[str, bytes, str]:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for _, name, data in files:
            z.writestr(name, data)
    return f"Documentacion - {_safe(ctx.name)}.zip", buf.getvalue(), "application/zip"
