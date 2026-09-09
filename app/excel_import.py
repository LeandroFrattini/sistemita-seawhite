"""Importa un line-up desde el Excel que manda la terminal.

Detecta cada bloque buscando una fila cuyo texto "Vessel type" aparezca en
alguna columna; a partir de ahi las columnas se leen en orden fijo. La celda
inmediatamente a la izquierda de "Vessel type" es, en la fila de encabezado,
el codigo de la terminal, y en las filas siguientes el nombre del barco.
Cada bloque termina en la primera fila sin nombre de barco.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

# offset relativo a la columna de "Vessel type"
FIELDS_AFTER = [
    "vessel_type", "imo", "eta", "etb", "etc", "operation",
    "quantity", "grade", "shipper", "destination", "local", "principal", "otras",
]


@dataclass
class ImportedCall:
    vessel_name: str = ""
    vessel_type: str = ""
    imo: str = ""
    eta: str = ""
    etb: str = ""
    etc: str = ""
    operation: str = ""
    quantity: str = ""
    grade: str = ""
    shipper: str = ""
    destination: str = ""
    local: str = ""
    principal: str = ""
    otras: str = ""


@dataclass
class ImportedTerminal:
    code: str
    calls: list[ImportedCall] = field(default_factory=list)


@dataclass
class ImportResult:
    port_name: str = ""
    terminals: list[ImportedTerminal] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def total_calls(self) -> int:
        return sum(len(t.calls) for t in self.terminals)


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm(value) -> str:
    return _cell_str(value).replace("\xa0", " ").strip()


def parse_lineup_xlsx(data: bytes) -> ImportResult:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    result = ImportResult()

    # nombre del puerto: primera fila con una sola celda de texto larga
    for row in rows[:6]:
        cells = [_norm(c) for c in row if _norm(c)]
        if len(cells) == 1 and len(cells[0]) > 6 and "vessel" not in cells[0].lower():
            result.port_name = cells[0]
            break

    r = 0
    while r < len(rows):
        row = rows[r]
        vt_col = _find_vessel_type_col(row)
        if vt_col is None:
            r += 1
            continue

        name_col = vt_col - 1
        code = _norm(row[name_col]) if name_col >= 0 else ""
        if not code:
            code = f"TERMINAL {len(result.terminals) + 1}"
        term = ImportedTerminal(code=code)

        r += 1
        while r < len(rows):
            drow = rows[r]
            name = _norm(drow[name_col]) if name_col < len(drow) else ""
            if not name:
                break
            if _find_vessel_type_col(drow) is not None:  # arranca otro bloque
                break
            call = ImportedCall(vessel_name=name)
            for i, fname in enumerate(FIELDS_AFTER):
                col = vt_col + i
                if col < len(drow):
                    setattr(call, fname, _norm(drow[col]))
            term.calls.append(call)
            r += 1

        if term.calls:
            result.terminals.append(term)
        else:
            result.warnings.append(f"Bloque '{code}' sin barcos, se omite.")

    if not result.terminals:
        result.warnings.append(
            "No se encontro ninguna tabla. Debe tener una fila con 'Vessel type' "
            "y el nombre de la terminal a la izquierda."
        )
    return result


def _find_vessel_type_col(row) -> int | None:
    for idx, value in enumerate(row):
        if _norm(value).lower() == "vessel type":
            return idx
    return None
