"""Proforma final (Administracion): estado del asistente, Excel descargable y lectura de ese mismo Excel.

El sistema NO guarda las proformas finales: el asistente vive en el navegador y se baja como Excel. Ese Excel
lleva una hoja oculta (_estado) con todos los datos, asi la pagina puede volver a leerlo para rehacer o modificar
una proforma vieja. Se guarda una "foto" de los valores usados (dolar, fee, cargas sociales), por eso una proforma
vieja se reabre con los valores de su momento y no con los de hoy.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

VERSION = 5
MARCA = "SEAWHITE-PROFORMA-FINAL"
HOJA_ESTADO = "_estado"
MAX_BYTES = 5 * 1024 * 1024
_CHUNK = 30000  # una celda de Excel admite 32767 caracteres
MAX_FEES = 40

DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]

PASOS = [
    (1, "Barco y dólar", "Elegí el barco y la cotización"),
    (2, "Agencia a facturar", "Fee y cargas sociales"),
    (3, "Datos del buque", "Horarios, medidas, TRN y calados"),
    (4, "Practicaje", "Entrada y salida, según calado"),
    (5, "Uso de muelle, fondeadero y vías navegables", "Días por bloques de 8 horas"),
    (6, "Amarre y desamarre", "Tarifa fija según el horario"),
    (7, "Habilitaciones de aduana", "Conceptos y valores"),
    (8, "Tally", "En base a lo ya completado"),
    (9, "Serenos", "Turnos y cargas sociales"),
    (10, "Migraciones y sanitary free pratique", "Según el voucher y el TRN"),
    (11, "Chequeo general", "Revisar y modificar todo"),
]


def tipo_dia(d: date, feriados: set[str], no_laborables: set[str] | frozenset = frozenset()) -> str:
    """FERIADO | SABADO | DOMINGO | NO_LABORABLE | SEMANA.

    El feriado manda sobre todo. El dia no laborable solo se distingue de un dia de semana: si cae sabado o
    domingo se toma como sabado o domingo (a confirmar con el usuario)."""
    if d.isoformat() in feriados:
        return "FERIADO"
    base = {5: "SABADO", 6: "DOMINGO"}.get(d.weekday(), "SEMANA")
    if base == "SEMANA" and d.isoformat() in no_laborables:
        return "NO_LABORABLE"
    return base


def con_cargas(costo: float, pct: float) -> float:
    """Costo con cargas sociales: el % se SUMA (165 sobre 100 da 265)."""
    return costo * (1 + pct / 100)


# Habilitaciones de aduana (customs overtime): todas las del voucher mas una libre. Cada una = cantidad x precio (USD 200 por defecto, editable).
ADUANA_CLAVES = [("entrance", "FOR ENTRANCE"), ("loading", "FOR LOADING"), ("permanence", "FOR PERMANENCE"),
                 ("bunkering", "FOR BUNKERING"), ("clearance", "FOR CLEARANCE"), ("store_list", "FOR STORE LIST ENDORSED"),
                 ("fresh_water", "FOR FRESH WATER"), ("provisions", "FOR PROVISIONS")]  # el mismo orden que el voucher
ADUANA_EXTRA = "extra"
ADUANA_USD = 200.0


def _num(v, tope: float) -> float:
    try:
        n = float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0
    return round(n, 4) if 0 <= n <= tope else 0.0


def _fecha(v) -> str:
    try:
        return date.fromisoformat(str(v)).isoformat()
    except (TypeError, ValueError):
        return ""


def _datetime(v) -> str:
    """Fecha y hora como la guarda 'Completar datos' (yyyy-mm-ddThh:mm), o vacio."""
    try:
        return datetime.strptime(str(v), "%Y-%m-%dT%H:%M").strftime("%Y-%m-%dT%H:%M")
    except (TypeError, ValueError):
        return ""


CARACTERES_INVALIDOS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff\ud800-\udfff]")  # no se pueden guardar en un Excel (XML 1.0)


def _texto(v, largo: int) -> str:
    return " ".join(CARACTERES_INVALIDOS.sub("", str(v or "")).split())[:largo]


def _turnos_limpios(raw) -> list[dict]:
    out = []
    for t in (raw if isinstance(raw, list) else [])[:120]:
        if isinstance(t, dict) and _fecha(t.get("fecha")) and str(t.get("turno")) in ("01-07", "07-13", "13-19", "19-01"):
            out.append({"fecha": _fecha(t["fecha"]), "turno": str(t["turno"]), "act": _texto(t.get("act"), 30).upper() or "LDG"})
    return out


def clean_buque(raw) -> dict:
    raw = raw if isinstance(raw, dict) else {}
    try:
        op_id = int(raw.get("operated_id")) if raw.get("operated_id") not in (None, "") else None
    except (TypeError, ValueError):
        op_id = None
    return {
        "operated_id": op_id,
        "nombre": _texto(raw.get("nombre"), 120),
        "imo": _texto(raw.get("imo"), 20),
        "sitio": _texto(raw.get("sitio"), 120),
        "amarre": _datetime(raw.get("amarre")),
        "inicio": _datetime(raw.get("inicio")),
        "fin": _datetime(raw.get("fin")),
        "zarpe": _datetime(raw.get("zarpe")),
        "cantidad": _texto(raw.get("cantidad"), 40),
        "carga": _texto(raw.get("carga"), 60),      # mercaderia (ej. CEBADA/ MALTA) y shipper: salen en la cabecera de la proforma
        "shipper": _texto(raw.get("shipper"), 60),
        # TRN y medidas: de ellos salen varios calculos de los pasos siguientes (practicaje usa eslora, manga, puntal y calados)
        "trn": _num(raw.get("trn"), 1_000_000),
        "trb": _num(raw.get("trb"), 1_000_000),
        "eslora": _num(raw.get("eslora"), 1000),
        "manga": _num(raw.get("manga"), 1000),
        "puntal": _num(raw.get("puntal"), 1000),
        "bandera": _texto(raw.get("bandera"), 80),
        "calado_entrada": _num(raw.get("calado_entrada"), 200),  # en pies, como la tabla de practicaje
        "calado_salida": _num(raw.get("calado_salida"), 200),
        # lo que se uso segun el voucher (foto): servicios y turnos de Head Tally
        "servicios": [_texto(x, 200) for x in (raw.get("servicios") if isinstance(raw.get("servicios"), list) else [])[:60] if _texto(x, 200)],
        "turnos": _turnos_limpios(raw.get("turnos")),
    }


def _aduana_items(raw, grupo_ajeno: bool = False) -> list[dict]:
    """Siempre las habilitaciones del voucher, en orden, mas la libre (el estado puede venir de un Excel ajeno o viejo).
    Cada una puede ir por cuenta del armador; un Excel viejo traia una sola casilla para todas (grupo_ajeno)."""
    por_clave = {}
    for it in (raw if isinstance(raw, list) else [])[:20]:
        if isinstance(it, dict) and it.get("key") in [k for k, _ in ADUANA_CLAVES] + [ADUANA_EXTRA]:
            por_clave.setdefault(it["key"], it)
    out = []
    for clave, texto in ADUANA_CLAVES + [(ADUANA_EXTRA, "")]:
        it = por_clave.get(clave, {})
        precio = _num(it.get("precio"), 1_000_000)
        out.append({
            "key": clave,
            "texto": _texto(it.get("texto"), 80) if clave == ADUANA_EXTRA else texto,
            "on": bool(it.get("on")),
            "cant": int(_num(it.get("cant"), 1000)),
            "precio": precio if "precio" in it else ADUANA_USD,
            "cuenta_ajena": bool(it.get("cuenta_ajena")) or grupo_ajeno,
            "omitido": bool(it.get("omitido")) and not (it.get("cuenta_ajena") or grupo_ajeno),
        })
    return out


def aduana_total(items: list[dict]) -> float:
    """Lo que se factura (lo que va por cuenta del armador no suma)."""
    return round(sum(i["cant"] * i["precio"] for i in items if i["on"] and not i["cuenta_ajena"] and not i["omitido"]), 2)


def clean_state(raw) -> dict:
    """Se queda solo con los campos conocidos y con valores validos (el estado puede venir de un Excel ajeno)."""
    raw = raw if isinstance(raw, dict) else {}
    dolar = raw.get("dolar") if isinstance(raw.get("dolar"), dict) else {}
    ag = raw.get("agencia") if isinstance(raw.get("agencia"), dict) else {}
    try:
        ag_id = int(ag.get("id")) if ag.get("id") not in (None, "") else None
    except (TypeError, ValueError):
        ag_id = None
    fees = []
    for f in (ag.get("fees") if isinstance(ag.get("fees"), list) else [])[:MAX_FEES]:
        if not isinstance(f, dict):
            continue
        texto = " ".join(str(f.get("texto") or "").split())[:200]
        if texto:
            fees.append({"texto": texto, "monto_usd": _num(f.get("monto_usd"), 10_000_000), "on": bool(f.get("on"))})
    # v1 tenia un solo "cargas_pct" y un solo "fee_usd": se leen como serenos y como un renglon de fee
    serenos = ag.get("serenos_pct", ag.get("cargas_pct"))
    if not fees and _num(ag.get("fee_usd"), 10_000_000):
        fees.append({"texto": "FEE", "monto_usd": _num(ag.get("fee_usd"), 10_000_000), "on": True})
    pr = raw.get("practicaje") if isinstance(raw.get("practicaje"), dict) else {}

    def mov(k):
        m = pr.get(k) if isinstance(pr.get(k), dict) else {}
        return {"usd": _num(m.get("usd"), 100_000_000), "manual": bool(m.get("manual")), "cuenta_ajena": bool(m.get("cuenta_ajena")),
                "omitido": bool(m.get("omitido")) and not m.get("cuenta_ajena")}

    def sub(k):
        return raw.get(k) if isinstance(raw.get(k), dict) else {}

    def concepto(d, extra=None):
        # cuenta_ajena = figura al pie (FOR YOUR ACCOUNT, no se factura); omitido = no figura en ningun lado de la proforma
        out = {"usd": _num(d.get("usd"), 100_000_000), "manual": bool(d.get("manual")), "cuenta_ajena": bool(d.get("cuenta_ajena")),
               "omitido": bool(d.get("omitido")) and not d.get("cuenta_ajena")}
        out.update(extra or {})
        return out

    mu, fo, vi = sub("muelle"), sub("fondeo"), sub("vias")
    la, ld, es, dr = sub("lds_amarre"), sub("lds_desamarre"), sub("escala"), sub("lds_draft")
    ad = sub("aduana")
    items = _aduana_items(ad.get("items"), bool(ad.get("cuenta_ajena")))
    from . import chequeo, migraciones, personal  # import tardio: usan tipo_dia / _num de este modulo

    return {
        "v": VERSION,
        "buque": clean_buque(raw.get("buque")),
        "practicaje": {"in": mov("in"), "out": mov("out")},
        # entrada/salida vacias = las del paso 3 (amarre y zarpe); "cuenta_ajena" = figura en FOR YOUR ACCOUNT (no se factura)
        "muelle": concepto(mu, {"entrada": _datetime(mu.get("entrada")), "salida": _datetime(mu.get("salida")),
                      # beneficio por amarre fuera de dia habil: se destilda con sin_beneficio; "perdonados" es la foto de los dias no cobrados
                      "sin_beneficio": bool(mu.get("sin_beneficio")), "perdonados": _num(mu.get("perdonados"), 100)}),
        "fondeo": concepto(fo, {"on": bool(fo.get("on")), "entrada": _datetime(fo.get("entrada")), "salida": _datetime(fo.get("salida"))}),
        "vias": concepto(vi, {"solo_bunker": bool(vi.get("solo_bunker")), "factor": _num(vi.get("factor"), 10) or 1.0}),
        # LDS: momento vacio = el amarre / el zarpe del paso 3; lanchas 0 = automatico (dos si la eslora supera 150 m)
        "lds_amarre": concepto(la, {"momento": _datetime(la.get("momento")), "lanchas": int(_num(la.get("lanchas"), 4))}),
        "lds_desamarre": concepto(ld, {"momento": _datetime(ld.get("momento")), "lanchas": int(_num(ld.get("lanchas"), 4))}),
        # escala (SHORE GANGWAY): on_set indica que alguien tildo o destildo la casilla (si no, manda lo que dice el voucher)
        "escala": concepto(es, {"on": bool(es.get("on")), "on_set": bool(es.get("on_set")), "dias": int(_num(es.get("dias"), 365))}),
        # lancha de draft (opcional): momento vacio = el fin de operacion del paso 3
        "lds_draft": concepto(dr, {"on": bool(dr.get("on")), "momento": _datetime(dr.get("momento"))}),
        # habilitaciones de aduana: cada renglon = cantidad x precio; el importe es la suma de los tildados (o el valor a mano)
        "aduana": concepto(ad, {"items": items}) | {"usd": aduana_total(items), "cuenta_ajena": False},
        # Tally y Serenos van en pesos: el estado guarda los renglones y los totales que se vieron en pantalla
        "migraciones": migraciones.clean(raw.get("migraciones")),
        "chequeo": chequeo.clean(raw.get("chequeo")),
        "tally": personal.clean_tally(raw.get("tally")),
        "serenos": personal.clean_serenos(raw.get("serenos")),
        "dolar": {
            "valor": _num(dolar.get("valor"), 1_000_000),
            "fecha": _fecha(dolar.get("fecha")),
            "origen": "BNA" if dolar.get("origen") == "BNA" else "MANUAL",
        },
        "agencia": {
            "id": ag_id,
            "nombre": str(ag.get("nombre") or "").strip()[:120],
            "serenos_pct": _num(serenos, 1000),
            "tally_pct": _num(ag.get("tally_pct"), 1000),
            "dolar_fee": _num(ag.get("dolar_fee"), 1_000_000),  # el fee se pasa a pesos con este dolar (0 = el mismo dolar del paso 1)
            "fees": fees,
        },
    }


def total_fee(state: dict) -> float:
    return round(sum(f["monto_usd"] for f in state["agencia"]["fees"] if f["on"]), 2)


def _nota_concepto(c: dict, entrada: str = "", salida: str = "") -> str:
    partes = []
    if entrada or salida:
        partes.append(f"{entrada.replace('T', ' ')} → {salida.replace('T', ' ')}")
    partes.append("modificado a mano" if c["manual"] else "calculado")
    if c["cuenta_ajena"]:
        partes.append("FOR YOUR ACCOUNT (no se factura)")
    return " · ".join(partes)


def _metros(pies: float) -> str:
    return f"{pies / 3.28084:.2f} m" if pies else "—"


def _hoja_chequeo(ws, r: int, ch: dict) -> int:
    """Chequeo general: renglones numerados en dolares y pesos, lo que va por cuenta del armador y los totales."""
    azul = PatternFill("solid", fgColor="1F67B0")
    for c, t in enumerate(("CHEQUEO GENERAL", "USD", "ARS"), start=1):
        x = ws.cell(r, c, t)
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = azul
    r += 1
    for l in ch["lineas"]:
        ws.cell(r, 1, f"{l['n']}  {l['texto']}")
        ws.cell(r, 2, l["usd"] if l["usd"] is not None else "")
        ws.cell(r, 3, l["ars"])
        r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 2, ch["total_usd"] if ch["total_usd"] is not None else "").font = Font(bold=True)
    ws.cell(r, 3, ch["total_ars"]).font = Font(bold=True)
    r += 2
    if ch["ajenas"] or ch["pie"]:
        ws.cell(r, 1, "FOR YOUR ACCOUNT").font = Font(bold=True)
        r += 1
        for l in ch["ajenas"]:
            ws.cell(r, 1, l["texto"])
            ws.cell(r, 2, l["usd"] if l["usd"] is not None else "")
            ws.cell(r, 3, "* sin impuestos")
            r += 1
        for t in ch["pie"]:
            ws.cell(r, 1, t)
            r += 1
    return r


def _nota_practicaje(m: dict) -> str:
    return ("modificado a mano" if m["manual"] else "según tabla de calado") + (" · FOR YOUR ACCOUNT (no se factura)" if m["cuenta_ajena"] else "")


USD_KEYS = ("mig_in", "mig_out")
ARS_KEYS = ("taxi_in", "taxi_out", "free_pratique", "taxi_sanidad")


def _mig_on(m: dict, claves) -> bool:
    return any(i["on"] for i in m["items"] if i["key"] in claves)


def _nota_migraciones(m: dict, claves) -> str:
    from . import migraciones

    textos = {k: t for k, t, *_ in migraciones.ITEMS}
    on = [i for i in m["items"] if i["key"] in claves and i["on"]]
    if not on:
        return "sin movimientos en el voucher"
    return " · ".join(textos[i["key"]] + (" EN BOYA" if i["boya"] else "") + (f" ({i['precio']:g} a mano)" if i["precio"] else "") for i in on)


def _usd_txt(ars: float, dolar: float) -> str:
    return f" · USD {ars / dolar:,.2f}" if ars and dolar else ""


def _nota_tally(t: dict, pct: float, dolar: float) -> str:
    if not t["turnos"]:
        return "sin turnos de Head Tally"
    from . import personal

    partes = [f"{r['cantidad']} × {personal.NIVELES_TALLY[r['nivel']]}" for r in t["resumen"] if r["cantidad"]]
    return (personal.CATEGORIAS_TALLY.get(t["categoria"], "SIN CATEGORÍA") + ": " + " + ".join(partes) + f" · cargas sociales {pct:g} % sumadas"
            + _usd_txt(t["total_ars"], dolar))


def _nota_serenos(s: dict, pct: float, dolar: float) -> str:
    if not s["turnos"]:
        return "sin turnos de serenos"
    from . import personal

    v = "sin viáticos" if s["viatico"] in ("", "NINGUNO") else personal.VIATICOS[s["viatico"]]
    return (f"{s['cantidad'] or len(s['turnos'])} turnos · {personal.CATEGORIAS_SERENOS.get(s['categoria'], 'categoría por turno').split(' (')[0]} · cargas sociales {pct:g} % sobre sueldos · {v}"
            + _usd_txt(s["total_ars"], dolar))


def _nota_aduana(a: dict) -> str:
    on = [i for i in a["items"] if i["on"]]
    if not on:
        return "sin habilitaciones"
    return " · ".join(f"{(i['texto'] or 'OTRO').upper()} {i['cant']} × {str(i['precio']).rstrip('0').rstrip('.')}" + (" (FOR YOUR ACCOUNT)" if i["cuenta_ajena"] else "") for i in on)


def _filas_pasos(s: dict) -> list[tuple]:
    """Lo cargado en cada paso, con su nota (calculado, modificado a mano, en boya, FOR YOUR ACCOUNT...): la foto de la proforma."""
    b = s["buque"]
    return [
        ("Buque", b["nombre"], (b["sitio"] + (" · IMO " + b["imo"] if b["imo"] else "")).strip(" ·")),
        ("Amarre", b["amarre"].replace("T", " "), ""),
        ("Inicio de operación", b["inicio"].replace("T", " "), ""),
        ("Fin de operación", b["fin"].replace("T", " "), ""),
        ("Zarpe", b["zarpe"].replace("T", " "), ""),
        ("Cantidad operada", b["cantidad"], ""),
        ("TRN", b["trn"] or "", "TRB " + str(b["trb"]) if b["trb"] else ""),
        ("Eslora / manga / puntal (m)", f"{b['eslora'] or '—'} / {b['manga'] or '—'} / {b['puntal'] or '—'}", b["bandera"] and ("Bandera " + b["bandera"]) or ""),
        ("Calado entrada / salida (pies)", f"{b['calado_entrada'] or '—'} / {b['calado_salida'] or '—'}",
         (f"en metros: {_metros(b['calado_entrada'])} / {_metros(b['calado_salida'])}" if b["calado_entrada"] or b["calado_salida"] else "")),
        ("Practicaje IN (USD)", s["practicaje"]["in"]["usd"] or "", _nota_practicaje(s["practicaje"]["in"])),
        ("Practicaje OUT (USD)", s["practicaje"]["out"]["usd"] or "", _nota_practicaje(s["practicaje"]["out"])),
        ("Uso de muelle (USD)", s["muelle"]["usd"] or "", _nota_concepto(s["muelle"], s["muelle"]["entrada"] or b["amarre"], s["muelle"]["salida"] or b["zarpe"])
         + (f" · beneficio por amarre fuera de día hábil: {str(s['muelle']['perdonados']).replace('.', ',')} días no cobrados" if s["muelle"]["perdonados"] and not s["muelle"]["sin_beneficio"] else "")),
        ("Uso de fondeadero (USD)", s["fondeo"]["usd"] or "" if s["fondeo"]["on"] else "—", _nota_concepto(s["fondeo"], s["fondeo"]["entrada"], s["fondeo"]["salida"]) if s["fondeo"]["on"] else "sin fondeo"),
        ("Vías navegables (USD)", s["vias"]["usd"] or "", _nota_concepto(s["vias"]) + (" · solo bunker: 20 % del TRN × " + str(s["vias"]["factor"]) if s["vias"]["solo_bunker"] else "")),
        ("Amarre LDS (USD)", s["lds_amarre"]["usd"] or "", _nota_concepto(s["lds_amarre"], s["lds_amarre"]["momento"] or b["amarre"])),
        ("Desamarre LDS (USD)", s["lds_desamarre"]["usd"] or "", _nota_concepto(s["lds_desamarre"], s["lds_desamarre"]["momento"] or b["zarpe"])),
        ("Escala LDS (USD)", s["escala"]["usd"] or "" if s["escala"]["on"] else "—", _nota_concepto(s["escala"]) if s["escala"]["on"] else "sin shore gangway"),
        ("Lancha de draft LDS (USD)", s["lds_draft"]["usd"] or "" if s["lds_draft"]["on"] else "—",
         _nota_concepto(s["lds_draft"], s["lds_draft"]["momento"] or b["fin"]) if s["lds_draft"]["on"] else "sin uso de lancha de draft"),
        ("Habilitaciones de aduana (USD)", s["aduana"]["usd"] or "" if any(i["on"] for i in s["aduana"]["items"]) else "—", _nota_aduana(s["aduana"])),
        ("Migraciones (USD)", s["migraciones"]["usd"] or "" if _mig_on(s["migraciones"], USD_KEYS) else "—", _nota_migraciones(s["migraciones"], USD_KEYS)),
        ("Free pratique y taxis (ARS)", s["migraciones"]["ars"] or "" if _mig_on(s["migraciones"], ARS_KEYS) else "—", _nota_migraciones(s["migraciones"], ARS_KEYS)
         + _usd_txt(s["migraciones"]["ars"], s["dolar"]["valor"])),
        ("Tally (ARS)", s["tally"]["total_ars"] or "" if s["tally"]["turnos"] else "—", _nota_tally(s["tally"], s["agencia"]["tally_pct"], s["dolar"]["valor"])),
        ("Serenos (ARS)", s["serenos"]["total_ars"] or "" if s["serenos"]["turnos"] else "—", _nota_serenos(s["serenos"], s["agencia"]["serenos_pct"], s["dolar"]["valor"])),
        ("Dólar (divisa BNA)", s["dolar"]["valor"] or "", "Cotización del " + (s["dolar"]["fecha"] or "—")
         + (" · traído del BNA" if s["dolar"]["origen"] == "BNA" else " · cargado a mano")),
        ("Agencia a facturar", s["agencia"]["nombre"], ""),
    ]


def _build_excel_resumen(state: dict, chequeo: dict | None = None) -> bytes:
    """Resumen en una sola hoja (sin el formato de las proformas): lo cargado en cada paso + hoja oculta con el estado. Es lo que sale sin base de datos."""
    s = clean_state(state)
    wb = Workbook()
    ws = wb.active
    ws.title = "Proforma"
    azul = PatternFill("solid", fgColor="1F67B0")
    ws["A1"] = "PROFORMA FINAL"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = azul
    ws.merge_cells("A1:C1")
    filas = _filas_pasos(s)
    r = 3
    for a, b, c in filas:
        ws.cell(r, 1, a).font = Font(bold=True)
        ws.cell(r, 2, b)
        ws.cell(r, 3, c)
        r += 1
    ws.cell(r, 1, "Fee aplicado").font = Font(bold=True)
    elegidos = [f for f in s["agencia"]["fees"] if f["on"]]
    if not elegidos:
        ws.cell(r, 2, "—")
        r += 1
    for f in elegidos:
        ws.cell(r, 2, f["monto_usd"])
        ws.cell(r, 3, f["texto"])
        r += 1
    ws.cell(r, 1, "Total fee (USD)").font = Font(bold=True)
    ws.cell(r, 2, total_fee(s))
    r += 1
    ws.cell(r, 1, "Cargas sociales Serenos (%)").font = Font(bold=True)
    ws.cell(r, 2, s["agencia"]["serenos_pct"])
    ws.cell(r, 3, "Se suman al costo: sobre 100 con 165 % da 265")
    r += 1
    ws.cell(r, 1, "Cargas sociales Head Tally (%)").font = Font(bold=True)
    ws.cell(r, 2, s["agencia"]["tally_pct"])
    r += 2
    if chequeo:
        r = _hoja_chequeo(ws, r, chequeo)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 52

    est = wb.create_sheet(HOJA_ESTADO)
    est["A1"] = f"{MARCA} v{VERSION}"
    texto = json.dumps(s, ensure_ascii=False, separators=(",", ":"))
    for i in range(0, len(texto), _CHUNK):
        est.cell(2 + i // _CHUNK, 1, texto[i:i + _CHUNK])
    est.sheet_state = "hidden"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _filas_resumen(s: dict) -> list[list]:
    """Hoja oculta `_resumen` del Excel final: la foto de cada paso con sus notas, como texto."""
    out = [["RESUMEN DE LA PROFORMA (foto de lo cargado en cada paso)"]]
    out += [[a, b, c] for a, b, c in _filas_pasos(s)]
    out += [["Fee aplicado", f["monto_usd"], f["texto"]] for f in s["agencia"]["fees"] if f["on"]] or [["Fee aplicado", "—", ""]]
    out += [["Total fee (USD)", total_fee(s), ""], ["Dólar del fee", s["agencia"]["dolar_fee"] or "igual al del paso 1", "El fee se pasa a pesos con este dólar"],
            ["Cargas sociales Serenos (%)", s["agencia"]["serenos_pct"], "Se suman al costo: sobre 100 con 165 % da 265"],
            ["Cargas sociales Head Tally (%)", s["agencia"]["tally_pct"], ""]]
    return out


def _filas_estado(s: dict) -> list[list]:
    texto = json.dumps(s, ensure_ascii=False, separators=(",", ":"))
    return [[f"{MARCA} v{VERSION}"]] + [[texto[i:i + _CHUNK]] for i in range(0, len(texto), _CHUNK)]


def _build_excel_final(state: dict, chequeo: dict | None, db) -> bytes:
    """El Excel final con el formato de las proformas (hoja de la agencia, SERENOS y TALLY), en pesos y en USD."""
    from . import chequeo as chq, muelle, personal
    from .calendario import conjuntos
    from .docgen import proforma_xlsx

    s = clean_state(state)
    ch = chequeo or chq.armar(db, s)
    b, dolar = s["buque"], s["dolar"]["valor"]
    feriados, _ = conjuntos(db)
    ref = date.fromisoformat(b["amarre"][:10]) if b["amarre"] else None
    tl, sr = s["tally"], s["serenos"]
    tally = personal.calcular_tally(db, categoria=tl["categoria"], turnos=tl["turnos"], precios=tl["precios"], pct=s["agencia"]["tally_pct"], dolar=dolar, ref=ref)
    serenos = personal.calcular_serenos(db, categoria=sr["categoria"], viatico=sr["viatico"], viatico_precio=sr["viatico_precio"], turnos=sr["turnos"] if sr["armado"] else None,
                                        amarre=b["amarre"], zarpe=b["zarpe"], pct=s["agencia"]["serenos_pct"], dolar=dolar, ref=ref)
    nombres = {(t["fecha"], t["ventana"]): t["nombre"] for t in sr["turnos"] if t.get("nombre")}
    for f in serenos["filas"]:
        f["nombre"] = nombres.get((f["fecha"], f["ventana"]), "")
    entrada, salida = s["muelle"]["entrada"] or b["amarre"], s["muelle"]["salida"] or b["zarpe"]
    libres = muelle.bloques_beneficio(db, entrada, b["inicio"]) if not s["muelle"]["sin_beneficio"] else frozenset()
    d = muelle.dias(entrada, salida, libres)
    dias = {"total": d["total"], "detalle": [x["subtotal"] for x in d["detalle"] if x["subtotal"] > 0]}
    return proforma_xlsx.generar(s, ch, tally, serenos, dias, feriados, _filas_resumen(s), _filas_estado(s))


def barco_clave(b: dict) -> str:
    """Identifica al barco para anotar lo pagado (el barco de Operaciones, o su nombre y dia de amarre si se escribio a mano)."""
    if b.get("operated_id"):
        return f"op{b['operated_id']}"
    return f"{(b.get('nombre') or '').upper()}|{(b.get('amarre') or '')[:10]}"[:160]


def es_barco_de_prueba(b: dict) -> bool:
    """El barco de prueba del paso 1 no deja nada anotado (son datos de ejemplo)."""
    return "BARCO DE PRUEBA" in (b.get("nombre") or "").upper()


def build_excel_serenos(state: dict, db, usuario: str = "") -> tuple[bytes, str]:
    """El Excel de SERENOS (recibos a cada sereno + RESULTADO + CUIL + WATCHMAN) con lo cargado en el paso 9. ValueError si falta algo.
    Anota lo cobrado por cada sereno en este barco (para su resumen), salvo en el barco de prueba."""
    from . import personal, serenos_pagos
    from .calendario import conjuntos
    from .docgen import proforma_xlsx, serenos_xlsx
    from sqlalchemy import select

    from .models import Sereno

    s = clean_state(state)
    b, dolar, sr = s["buque"], s["dolar"]["valor"], s["serenos"]
    if sr["omitido"]:
        raise ValueError("Los serenos están marcados como que no figuran en la proforma (paso 9): no hay recibos para armar.")
    ref = date.fromisoformat(b["amarre"][:10]) if b["amarre"] else None
    serenos = personal.calcular_serenos(db, categoria=sr["categoria"], viatico=sr["viatico"], viatico_precio=sr["viatico_precio"], turnos=sr["turnos"] if sr["armado"] else None,
                                        amarre=b["amarre"], zarpe=b["zarpe"], pct=s["agencia"]["serenos_pct"], dolar=dolar, ref=ref)
    nombres = {(t["fecha"], t["ventana"]): t["nombre"] for t in sr["turnos"] if t.get("nombre")}
    for f in serenos["filas"]:
        f["nombre"] = nombres.get((f["fecha"], f["ventana"]), "")
    ref_tabla = ref or personal._d(b["amarre"]) or next((personal._d(f["fecha"]) for f in serenos["filas"]), None)  # la misma fecha con la que se calcularon los turnos
    base, _ = personal.tabla(db, ref_tabla)
    feriados, _ = conjuntos(db)
    padron = [(x.nombre, x.cuil) for x in db.scalars(select(Sereno).order_by(Sereno.nombre))]
    titulo = proforma_xlsx.titulo_buque(b["nombre"])
    datos, recibos = serenos_xlsx.generar(buque=titulo, sitio=b["sitio"], serenos=serenos, base=base, padron=padron, pct=s["agencia"]["serenos_pct"], dolar=dolar or 0.0,
                                          feriados=feriados, hoy=date.today())
    if not es_barco_de_prueba(b):
        serenos_pagos.registrar(db, barco_clave=barco_clave(b), buque=titulo, recibos=recibos, usuario=usuario)
    limpio = re.sub(r"[^A-Za-z0-9 ._-]", "", titulo).strip() or "buque"
    return datos, f"SERENOS - {limpio} - {date.today().strftime('%d.%m.%Y')}.xlsx"


def build_excel_encargados(state: dict, db, usuario: str = "") -> tuple[bytes, str]:
    """El Excel de ENCARGADOS (declaracion jurada del Head Tally) con lo cargado en el paso 8. ValueError si falta algo.
    Anota los turnos de este barco para llevar el tope mensual de aportes, salvo en el barco de prueba."""
    from sqlalchemy import select

    from . import encargados as enc, personal
    from .docgen import encargados_xlsx, proforma_xlsx
    from .models import Encargado
    from .proformador_calc import get_param

    s = clean_state(state)
    b, dolar, tl = s["buque"], s["dolar"]["valor"], s["tally"]
    if tl["omitido"]:
        raise ValueError("El Head Tally está marcado como que no figura en la proforma (paso 8): no hay encargados para armar.")
    ref = date.fromisoformat(b["amarre"][:10]) if b["amarre"] else None
    t = personal.calcular_tally(db, categoria=tl["categoria"], turnos=tl["turnos"], precios=tl["precios"], pct=0, dolar=dolar, ref=ref)
    nombres = {(x["fecha"], x["turno"]): x["nombre"] for x in tl["turnos"]}
    filas = [{**f, "nombre": nombres.get((f["fecha"], f["turno"]), "")} for f in t["filas"]]
    ref_tabla = ref or next((personal._d(f["fecha"]) for f in filas), None)  # la misma fecha con la que se calcularon los turnos
    liquidables = [f for f in filas if not enc.es_berthing(f.get("act"))]  # el BERTHING esta en la proforma pero no se liquida al encargado
    if filas and not liquidables:
        raise ValueError("Todos los turnos del Head Tally son BERTHING: no hay nada para liquidar a los encargados.")
    filas = liquidables
    base, _ = personal.tabla(db, ref_tabla)
    lista = list(db.scalars(select(Encargado).order_by(Encargado.nombre)))
    cuil_de = {e.nombre: e.cuil for e in lista}
    for f in filas:
        if f["nombre"] and f["nombre"] not in cuil_de:
            raise ValueError(f"{f['nombre']} ya no está en la lista de encargados (Administración → Encargados): elegí otro encargado en el paso 8.")
    calc = enc.calcular(db, barco_clave=barco_clave(b), filas=filas, categoria=t["categoria"], base=base, padron=[(e.id, e.nombre) for e in lista])
    titulo = proforma_xlsx.titulo_buque(b["nombre"])
    operacion = "DESCARGA" if any(str(f.get("act") or "").upper().startswith(("DISCH", "DESC")) for f in filas) else "CARGA"
    plano = re.sub(r"^(MV|M/V|M\.V\.)\s+", "", titulo)  # en la declaracion el barco va sin el MV, como en el libro
    datos = encargados_xlsx.generar(buque=plano, operacion=operacion, carga=(b["carga"] or "").upper(), sitio=(b["sitio"] or "").upper(), calc=calc, cuil_de=cuil_de,
                                    viatico_dia=get_param(db, "viatico_encargado_ars", 22933.49), hoy=date.today())
    if not es_barco_de_prueba(b):
        enc.registrar(db, barco_clave=barco_clave(b), buque=titulo, filas=calc["filas"], usuario=usuario)
    limpio = re.sub(r"[^A-Za-z0-9 ._-]", "", titulo).strip() or "buque"
    return datos, f"ENCARGADOS - {limpio} - {date.today().strftime('%d.%m.%Y')}.xlsx"


def build_excel(state: dict, chequeo: dict | None = None, db=None) -> bytes:
    """Con base de datos: el Excel final con el formato de las proformas. Sin ella: el resumen simple (ver _build_excel_resumen)."""
    if db is None:
        return _build_excel_resumen(state, chequeo)
    return _build_excel_final(state, chequeo, db)


def read_excel(data: bytes) -> dict:
    """Lee el estado de un Excel bajado por este sistema. ValueError si no lo es."""
    if len(data) > MAX_BYTES:
        raise ValueError("El archivo es demasiado grande para ser una proforma del sistema.")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise ValueError("No se pudo abrir el archivo. Tiene que ser un Excel (.xlsx) bajado desde este sistema.")
    if HOJA_ESTADO not in wb.sheetnames:
        raise ValueError("Este Excel no fue generado por el sistema (falta la hoja con los datos de la proforma).")
    ws = wb[HOJA_ESTADO]
    if not str(ws["A1"].value or "").startswith(MARCA):
        raise ValueError("Este Excel no fue generado por el sistema.")
    partes = []
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0] is None:
            break
        partes.append(str(row[0]))
    try:
        return clean_state(json.loads("".join(partes)))
    except ValueError:
        raise ValueError("Los datos guardados en el Excel están dañados o fueron modificados.")
