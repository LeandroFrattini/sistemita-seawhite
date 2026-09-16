import io
import re
import urllib.request

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import admin_required, current_user, pda_admin_required, verify_password
from ..config import BASE_DIR
from ..database import get_db
from ..models import (
    Proforma,
    ProformaBunker,
    ProformaBunkerBoya,
    ProformaBunkerLinea,
    ProformaCoefTramo,
    ProformaConceptoFijo,
    ProformaLinea,
    ProformaParametro,
    ProformaPilotageTramo,
    ProformaTarifaTurno,
    ProformaTugTarifa,
    User,
)
from ..proformador_calc import calcular_bunker, calcular_fc, calcular_pilotage, calcular_proforma
from ..templating import templates

router = APIRouter()

TIPOS_BUQUE = ["Bulk Carrier", "LPG", "Tanker"]
TIPOS_OPERACION = ["Carga", "Descarga", "Bunker"]
TIPOS_CARGA_TALLY = [
    ("ACEITE", "Aceite"), ("CEREAL", "Cereal"), ("BOLSONES", "Bolsones"),
    ("FERTILIZANTE", "Fertilizante / Baritina a granel"),
]
CATEGORIAS_WATCHMEN = [
    ("NORMAL", "Normal - Aceite"), ("INSALUBRE", "Insalubre - Cereal/Fertilizante"),
    ("PELIGROSO", "Peligroso - Postas"),
]
DIAS_TIPO_LABEL = [("SEMANA", "Día de semana"), ("SABADO", "Sábado"), ("DOMINGO_FERIADO", "Domingo/Feriado")]
BOYAS = [("BOYA_3", "Boya 3"), ("BOYA_11", "Boya 11"), ("BOYA_17", "Boya 17")]

LOGO_PATH = BASE_DIR / "app" / "static" / "img" / "logo-sw-emblem.png"


def _insertar_logo(ws, cell="C5", width=70, height=83, offset_x_px=6, offset_y_px=16):
    if not LOGO_PATH.exists():
        return
    img = XLImage(str(LOGO_PATH))
    img.width = width
    img.height = height
    row, col = coordinate_to_tuple(cell)
    marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(offset_x_px), row=row - 1, rowOff=pixels_to_EMU(offset_y_px))
    img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)))
    ws.add_image(img)


@router.get("/proformador", response_class=HTMLResponse)
def listar_proformas(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    proformas = db.scalars(select(Proforma).order_by(Proforma.creado_en.desc()).limit(100)).all()
    return templates.TemplateResponse(request, "proformador/lista.html", {"user": user, "proformas": proformas})


@router.get("/proformador/nuevo", response_class=HTMLResponse)
def nuevo_wizard(request: Request, user: User = Depends(current_user)):
    return templates.TemplateResponse(request, "proformador/wizard.html", {
        "user": user,
        "tipos_buque": TIPOS_BUQUE,
        "tipos_operacion": TIPOS_OPERACION,
        "tipos_carga_tally": TIPOS_CARGA_TALLY,
        "categorias_watchmen": CATEGORIAS_WATCHMEN,
        "dias_tipo": DIAS_TIPO_LABEL,
    })


@router.get("/proformador/pilotage/probar")
def probar_pilotage(uf: float = 0, calado: float = 0, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Calculadora suelta para probar la formula de pilotaje contra el
    tarifario (cargando el coeficiente/UF y el calado a mano), sin tener
    que armar una proforma completa."""
    valor = calcular_pilotage(db, uf, calado)
    return {"ok": True, "valor": round(valor)}


@router.get("/proformador/cotizacion")
def cotizacion_bna(user: User = Depends(current_user)):
    """Trae la cotizacion divisa (compra/venta) del Dolar U.S.A publicada por
    el Banco Nacion. Si el sitio no responde o cambia de formato, se devuelve
    un error y el campo queda editable a mano igual."""
    try:
        req = urllib.request.Request("https://www.bna.com.ar/Personas", headers={"User-Agent": "Mozilla/5.0"})
        html = urllib.request.urlopen(req, timeout=10).read().decode("utf-8", errors="ignore")
        bloque = re.search(r'id="divisas".*?</table>', html, re.S)
        m = re.search(
            r"Dolar U\.S\.A</td>\s*<td>([\d.,]+)</td>\s*<td>([\d.,]+)</td>",
            bloque.group(0) if bloque else "",
        )
        if not m:
            raise ValueError("no se encontro la cotizacion en la pagina del BNA")
        compra = float(m.group(1).replace(",", "."))
        venta = float(m.group(2).replace(",", "."))
        return {"ok": True, "compra": compra, "venta": venta}
    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)}, status_code=502)


# --- BUNKER (Boya 3 / Boya 11 / Boya 17) ---------------------------------
# OJO: estas rutas literales van ANTES de "/proformador/{proforma_id}" mas
# abajo, para que "/proformador/bunker..." no matchee ahi primero.

@router.get("/proformador/bunker", response_class=HTMLResponse)
def listar_bunkers(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    proformas = db.scalars(select(ProformaBunker).order_by(ProformaBunker.creado_en.desc()).limit(100)).all()
    return templates.TemplateResponse(request, "proformador/bunker_lista.html", {"user": user, "proformas": proformas})


@router.get("/proformador/bunker/nuevo", response_class=HTMLResponse)
def nuevo_bunker_wizard(request: Request, user: User = Depends(current_user)):
    return templates.TemplateResponse(request, "proformador/bunker_wizard.html", {"user": user, "boyas": BOYAS})


def _datos_bunker_from_form(form) -> dict:
    def f(key, default=0.0):
        v = form.get(key)
        try:
            return float(v) if v not in (None, "") else default
        except ValueError:
            return default

    eslora, manga, puntal = f("eslora"), f("manga"), f("puntal")
    return {
        "dolar_venta": f("dolar_venta"),
        "cliente": (form.get("cliente") or "").strip(),
        "nombre_buque": (form.get("nombre_buque") or "MV TBN").strip() or "MV TBN",
        "boya": form.get("boya") or "BOYA_11",
        "eslora": eslora, "manga": manga, "puntal": puntal,
        "fc": calcular_fc(eslora, manga, puntal),
        "trn": f("trn"),
        "calado_entrada": f("calado_entrada"),
        "calado_salida": f("calado_salida"),
        "dias_estadia": f("dias_estadia", 1),
        "cantidad_barcazas": int(f("cantidad_barcazas", 1)),
        "turnos_customs_clearance": f("turnos_customs_clearance", 2),
        "turnos_customs_bunker_control": f("turnos_customs_bunker_control", 4),
        "usa_boat_surveyor": form.get("usa_boat_surveyor") == "on",
        "horas_boat_surveyor": f("horas_boat_surveyor"),
        "turnos_sipa": f("turnos_sipa", 6),
        "procede_exterior": form.get("procede_exterior") == "on",
        "destino_exterior": form.get("destino_exterior") == "on",
    }


@router.post("/proformador/bunker/calcular")
async def calcular_bunker_route(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    form = await request.form()
    datos = _datos_bunker_from_form(form)
    lineas = calcular_bunker(db, datos)
    total = sum(l["monto_usd"] for l in lineas if not l["informativo"])
    return templates.TemplateResponse(request, "proformador/_bunker_preview.html", {
        "user": user, "datos": datos, "lineas": lineas, "total": round(total, 2), "boyas": dict(BOYAS),
    })


@router.post("/proformador/bunker/guardar")
async def guardar_bunker(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    form = await request.form()
    datos = _datos_bunker_from_form(form)
    lineas = calcular_bunker(db, datos)
    total = sum(l["monto_usd"] for l in lineas if not l["informativo"])

    p = ProformaBunker(
        dolar_venta=datos["dolar_venta"], cliente=datos["cliente"], nombre_buque=datos["nombre_buque"],
        boya=datos["boya"], eslora=datos["eslora"], manga=datos["manga"], puntal=datos["puntal"],
        fc=datos["fc"], trn=datos["trn"], calado_entrada=datos["calado_entrada"], calado_salida=datos["calado_salida"],
        dias_estadia=datos["dias_estadia"], cantidad_barcazas=datos["cantidad_barcazas"],
        turnos_customs_clearance=datos["turnos_customs_clearance"],
        turnos_customs_bunker_control=datos["turnos_customs_bunker_control"],
        usa_boat_surveyor=datos["usa_boat_surveyor"], horas_boat_surveyor=datos["horas_boat_surveyor"],
        turnos_sipa=datos["turnos_sipa"], procede_exterior=datos["procede_exterior"],
        destino_exterior=datos["destino_exterior"], total_usd=round(total, 2), creado_por=user.username,
    )
    db.add(p)
    db.flush()
    for i, l in enumerate(lineas):
        db.add(ProformaBunkerLinea(
            proforma_id=p.id, concepto=l["concepto"], monto_usd=l["monto_usd"],
            observacion=l["observacion"], informativo=l["informativo"], orden=i,
        ))
    db.commit()
    return RedirectResponse(url=f"/proformador/bunker/{p.id}", status_code=302)


@router.post("/proformador/bunker/{proforma_id}/eliminar")
def eliminar_bunker(proforma_id: int, db: Session = Depends(get_db), user: User = Depends(admin_required)):
    p = db.get(ProformaBunker, proforma_id)
    if p:
        db.delete(p)
        db.commit()
    return RedirectResponse(url="/proformador/bunker", status_code=302)


@router.get("/proformador/bunker/{proforma_id}", response_class=HTMLResponse)
def ver_bunker(request: Request, proforma_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p = db.get(ProformaBunker, proforma_id)
    if not p:
        return RedirectResponse(url="/proformador/bunker", status_code=302)
    total = sum(l.monto_usd for l in p.items if not l.informativo)
    return templates.TemplateResponse(request, "proformador/bunker_ver.html", {
        "user": user, "p": p, "total": round(total, 2), "boyas": dict(BOYAS),
    })


@router.get("/proformador/bunker/{proforma_id}/export.xlsx")
def exportar_bunker_xlsx(proforma_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p = db.get(ProformaBunker, proforma_id)
    if not p:
        return RedirectResponse(url="/proformador/bunker", status_code=302)

    BLUE = "FF1B5FA8"
    GREY = "FFF4F6FA"
    WHITE_BOLD = Font(bold=True, color="FFFFFFFF", size=11)
    TITLE_FONT = Font(bold=True, color="FFFFFFFF", size=15)
    BOLD = Font(bold=True)
    THIN = Side(style="thin", color="FFB7C6DC")
    BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def fill(color):
        return PatternFill("solid", fgColor=color)

    wb = Workbook()
    ws = wb.active
    ws.title = "PDA Bunker"
    ws.sheet_view.showGridLines = False
    last_col = 3

    ws.merge_cells("A1:C2")
    c = ws["A1"]
    c.value = "SEA WHITE S.A."
    c.font = TITLE_FONT
    c.fill = fill(BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for r in range(1, 3):
        for col in range(1, last_col + 1):
            ws.cell(row=r, column=col).fill = fill(BLUE)
    _insertar_logo(ws)

    ws["A3"] = "Facundo Zuviria 401 - Bahia Blanca, Argentina"
    ws["A4"] = "Mail: operations@seawhite.com.ar"

    ws["A6"] = "TO Messrs"
    ws["A6"].font = BOLD
    ws["B6"] = p.cliente or ""
    ws["A7"] = "Vessel"
    ws["A7"].font = BOLD
    ws["B7"] = p.nombre_buque or ""
    ws["B7"].font = BOLD
    ws["A8"] = "Bunkering only -- " + dict(BOYAS).get(p.boya, p.boya)
    ws["A8"].font = BOLD

    ws.merge_cells("A10:C10")
    hdr = ws["A10"]
    hdr.value = "VESSEL PARTICULARS"
    hdr.font = WHITE_BOLD
    hdr.fill = fill(BLUE)
    hdr.alignment = Alignment(horizontal="left", indent=1)

    caract = [
        ("NRT", p.trn), ("LOA", p.eslora), ("Beam", p.manga), ("Depth", p.puntal),
        ("Draft IN (ft)", p.calado_entrada), ("Draft OUT (ft)", p.calado_salida), ("FC", p.fc),
    ]
    caract = [(label, value) for label, value in caract if value]
    row = 11
    for label, value in caract:
        a, b = ws.cell(row=row, column=1, value=label), ws.cell(row=row, column=2, value=value)
        a.font = BOLD
        a.fill = fill(GREY)
        a.border = BOX
        b.border = BOX
        ws.cell(row=row, column=3).border = BOX
        row += 1

    row += 1
    headers = ["Concept", "Value (USD)", "Remark"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=1 + i, value=h)
        cell.font = WHITE_BOLD
        cell.fill = fill(BLUE)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="left", indent=1)
    row += 1
    for item in p.items:
        a = ws.cell(row=row, column=1, value=item.concepto)
        b = ws.cell(row=row, column=2, value=item.monto_usd)
        cc = ws.cell(row=row, column=3, value=item.observacion or "")
        b.number_format = "#,##0"
        for cell in (a, b, cc):
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=(cell is cc))
        row += 1

    total = sum(i.monto_usd for i in p.items if not i.informativo)
    total_label = ws.cell(row=row, column=1, value="TOTAL USD")
    total_value = ws.cell(row=row, column=2, value=round(total))
    total_label.font = WHITE_BOLD
    total_value.font = WHITE_BOLD
    total_value.number_format = "#,##0"
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(BLUE)
        ws.cell(row=row, column=col).border = BOX

    ws.row_dimensions[1].height = 22
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"PDA Bunker {p.nombre_buque or 'buque'}.xlsx".replace("/", "-")
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _datos_from_form(form) -> dict:
    def f(key, default=0.0):
        v = form.get(key)
        try:
            return float(v) if v not in (None, "") else default
        except ValueError:
            return default

    eslora, manga, puntal = f("eslora"), f("manga"), f("puntal")
    return {
        "dolar_venta": f("dolar_venta"),
        "cliente": (form.get("cliente") or "").strip(),
        "tipo_buque": form.get("tipo_buque") or TIPOS_BUQUE[0],
        "tipo_operacion": form.get("tipo_operacion") or TIPOS_OPERACION[0],
        "nombre_buque": (form.get("nombre_buque") or "MV TBN").strip() or "MV TBN",
        "eslora": eslora, "manga": manga, "puntal": puntal,
        "fc": calcular_fc(eslora, manga, puntal),
        "trn": f("trn"),
        "calado_entrada": f("calado_entrada"),
        "calado_salida": f("calado_salida"),
        "cantidad": f("cantidad"),
        "dias_muelle": f("dias_muelle"),
        "dias_fondeo": f("dias_fondeo"),
        "remolques_in": int(f("remolques_in")),
        "remolques_out": int(f("remolques_out")),
        "turnos": f("turnos"),
        "tipo_carga": form.get("tipo_carga") or "ACEITE",
        "categoria_watchmen": form.get("categoria_watchmen") or "NORMAL",
        "dia_tipo": form.get("dia_tipo") or "SEMANA",
        "procede_exterior": form.get("procede_exterior") == "on",
        "destino_exterior": form.get("destino_exterior") == "on",
        "immigration_in_boya": form.get("immigration_in_boya") == "on",
        "immigration_out_boya": form.get("immigration_out_boya") == "on",
    }


@router.post("/proformador/calcular")
async def calcular(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    form = await request.form()
    datos = _datos_from_form(form)
    lineas = calcular_proforma(db, datos)
    total = sum(l["monto_usd"] for l in lineas if not l["informativo"])
    return templates.TemplateResponse(request, "proformador/_preview.html", {
        "user": user, "datos": datos, "lineas": lineas, "total": round(total, 2),
    })


@router.post("/proformador/guardar")
async def guardar(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    form = await request.form()
    datos = _datos_from_form(form)
    lineas = calcular_proforma(db, datos)
    total = sum(l["monto_usd"] for l in lineas if not l["informativo"])

    p = Proforma(
        dolar_venta=datos["dolar_venta"], cliente=datos["cliente"], tipo_buque=datos["tipo_buque"],
        tipo_operacion=datos["tipo_operacion"], nombre_buque=datos["nombre_buque"],
        eslora=datos["eslora"], manga=datos["manga"], puntal=datos["puntal"], fc=datos["fc"], trn=datos["trn"],
        calado_entrada=datos["calado_entrada"], calado_salida=datos["calado_salida"],
        cantidad=datos["cantidad"], dias_muelle=datos["dias_muelle"], dias_fondeo=datos["dias_fondeo"],
        remolques_in=datos["remolques_in"], remolques_out=datos["remolques_out"], turnos=datos["turnos"],
        tipo_carga=datos["tipo_carga"], categoria_watchmen=datos["categoria_watchmen"], dia_tipo=datos["dia_tipo"],
        procede_exterior=datos["procede_exterior"], destino_exterior=datos["destino_exterior"],
        immigration_in_boya=datos["immigration_in_boya"], immigration_out_boya=datos["immigration_out_boya"],
        total_usd=round(total, 2), creado_por=user.username,
    )
    db.add(p)
    db.flush()
    for i, l in enumerate(lineas):
        db.add(ProformaLinea(
            proforma_id=p.id, concepto=l["concepto"], monto_usd=l["monto_usd"],
            observacion=l["observacion"], informativo=l["informativo"], orden=i,
        ))
    db.commit()
    return RedirectResponse(url=f"/proformador/{p.id}", status_code=302)


@router.post("/proformador/{proforma_id}/eliminar")
def eliminar_proforma(proforma_id: int, db: Session = Depends(get_db), user: User = Depends(admin_required)):
    p = db.get(Proforma, proforma_id)
    if p:
        db.delete(p)
        db.commit()
    return RedirectResponse(url="/proformador", status_code=302)


@router.get("/proformador/{proforma_id}", response_class=HTMLResponse)
def ver_proforma(request: Request, proforma_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p = db.get(Proforma, proforma_id)
    if not p:
        return RedirectResponse(url="/proformador", status_code=302)
    total = sum(l.monto_usd for l in p.items if not l.informativo)
    return templates.TemplateResponse(request, "proformador/ver.html", {"user": user, "p": p, "total": round(total, 2)})


@router.get("/proformador/{proforma_id}/export.xlsx")
def exportar_xlsx(proforma_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    p = db.get(Proforma, proforma_id)
    if not p:
        return RedirectResponse(url="/proformador", status_code=302)

    BLUE = "FF1B5FA8"
    BLUE_LIGHT = "FFCFE0F5"
    GREY = "FFF4F6FA"
    WHITE_BOLD = Font(bold=True, color="FFFFFFFF", size=11)
    TITLE_FONT = Font(bold=True, color="FFFFFFFF", size=15)
    BOLD = Font(bold=True)
    ITALIC_MUTED = Font(italic=True, color="FF64758A")
    THIN = Side(style="thin", color="FFB7C6DC")
    BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def fill(color):
        return PatternFill("solid", fgColor=color)

    wb = Workbook()
    ws = wb.active
    ws.title = "PDA"
    ws.sheet_view.showGridLines = False
    last_col = 3  # A..C

    # --- encabezado con la marca ---
    ws.merge_cells("A1:C2")
    c = ws["A1"]
    c.value = "SEA WHITE S.A."
    c.font = TITLE_FONT
    c.fill = fill(BLUE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for r in range(1, 3):
        for col in range(1, last_col + 1):
            ws.cell(row=r, column=col).fill = fill(BLUE)
    _insertar_logo(ws)

    ws["A3"] = "Facundo Zuviria 401 - Bahia Blanca, Argentina"
    ws["A4"] = "Mail: operations@seawhite.com.ar"
    ws["A3"].font = ITALIC_MUTED
    ws["A4"].font = ITALIC_MUTED

    ws["A6"] = "TO Messrs"
    ws["A6"].font = BOLD
    ws["B6"] = p.cliente or ""
    ws["A7"] = "Vessel"
    ws["A7"].font = BOLD
    ws["B7"] = p.nombre_buque or ""
    ws["B7"].font = BOLD

    # --- caracteristicas del buque ---
    ws.merge_cells("A9:C9")
    hdr = ws["A9"]
    hdr.value = "VESSEL PARTICULARS"
    hdr.font = WHITE_BOLD
    hdr.fill = fill(BLUE)
    hdr.alignment = Alignment(horizontal="left", indent=1)

    caract = [
        ("NRT", p.trn), ("Cargo (tn)", p.cantidad), ("LOA", p.eslora), ("Beam", p.manga),
        ("Depth", p.puntal), ("Draft IN (ft)", p.calado_entrada), ("Draft OUT (ft)", p.calado_salida), ("FC", p.fc),
        ("Tugs", (p.remolques_in or 0) + (p.remolques_out or 0)),
        ("Days alongside", p.dias_muelle), ("Shifts", p.turnos),
    ]
    caract = [(label, value) for label, value in caract if value]
    row = 10
    for label, value in caract:
        a, b = ws.cell(row=row, column=1, value=label), ws.cell(row=row, column=2, value=value)
        a.font = BOLD
        a.fill = fill(GREY)
        a.border = BOX
        b.border = BOX
        ws.cell(row=row, column=3).border = BOX
        row += 1

    # --- conceptos ---
    row += 1
    headers = ["Concept", "Value (USD)", "Remark"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=1 + i, value=h)
        cell.font = WHITE_BOLD
        cell.fill = fill(BLUE)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="left", indent=1)
    row += 1
    for item in p.items:
        a = ws.cell(row=row, column=1, value=item.concepto)
        b = ws.cell(row=row, column=2, value=item.monto_usd)
        cc = ws.cell(row=row, column=3, value=item.observacion or "")
        b.number_format = "#,##0"
        for cell in (a, b, cc):
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=(cell is cc))
        if item.informativo:
            for cell in (a, b, cc):
                cell.font = ITALIC_MUTED
            a.value = f"{item.concepto} (informativo, no suma al total)"
        row += 1
    last_item_row = row - 1

    total = sum(i.monto_usd for i in p.items if not i.informativo)
    total_label = ws.cell(row=row, column=1, value="TOTAL USD")
    total_value = ws.cell(row=row, column=2, value=round(total, 2))
    total_label.font = WHITE_BOLD
    total_value.font = WHITE_BOLD
    total_value.number_format = "#,##0"
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(BLUE)
        ws.cell(row=row, column=col).border = BOX

    ws.row_dimensions[1].height = 22
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 55

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"PDA {p.nombre_buque or 'buque'}.xlsx".replace("/", "-")
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# --- FORMULAS: visibles a todos, editables solo admin o is_pda_admin ----- #

@router.get("/proformador/config/formulas", response_class=HTMLResponse)
def ver_formulas(request: Request, db: Session = Depends(get_db), user: User = Depends(current_user)):
    puede_editar = user.is_admin or user.is_pda_admin
    return templates.TemplateResponse(request, "proformador/formulas.html", {
        "user": user,
        "puede_editar": puede_editar,
        "error": request.query_params.get("error"),
        "parametros": db.scalars(select(ProformaParametro).order_by(ProformaParametro.orden)).all(),
        "coef_tramos": db.scalars(select(ProformaCoefTramo).order_by(ProformaCoefTramo.orden)).all(),
        "tug_tarifas": db.scalars(select(ProformaTugTarifa).order_by(ProformaTugTarifa.orden)).all(),
        "pilotage_tramos": db.scalars(select(ProformaPilotageTramo).order_by(ProformaPilotageTramo.orden)).all(),
        "tarifas_turno": db.scalars(select(ProformaTarifaTurno).order_by(ProformaTarifaTurno.orden)).all(),
        "conceptos_fijos": db.scalars(select(ProformaConceptoFijo).order_by(ProformaConceptoFijo.orden)).all(),
        "bunker_boyas": db.scalars(select(ProformaBunkerBoya).order_by(ProformaBunkerBoya.orden)).all(),
    })


@router.post("/proformador/config/formulas/guardar")
async def guardar_formulas(request: Request, db: Session = Depends(get_db), user: User = Depends(pda_admin_required)):
    form = await request.form()

    password = form.get("confirm_password", "")
    if not password or not verify_password(password, user.password_hash):
        return RedirectResponse(url="/proformador/config/formulas?error=password", status_code=302)

    for p in db.scalars(select(ProformaParametro)):
        v = form.get(f"param_{p.id}")
        if v not in (None, ""):
            try:
                p.valor = float(v)
            except ValueError:
                pass

    for t in db.scalars(select(ProformaCoefTramo)):
        hasta = form.get(f"tramo_hasta_{t.id}")
        coef = form.get(f"tramo_coef_{t.id}")
        t.hasta_toneladas = float(hasta) if hasta not in (None, "") else None
        if coef not in (None, ""):
            try:
                t.coeficiente = float(coef)
            except ValueError:
                pass

    for tg in db.scalars(select(ProformaTugTarifa)):
        hasta = form.get(f"tug_hasta_{tg.id}")
        valor = form.get(f"tug_valor_{tg.id}")
        tg.hasta_loa = float(hasta) if hasta not in (None, "") else None
        if valor not in (None, ""):
            try:
                tg.valor_usd = float(valor)
            except ValueError:
                pass

    for pt in db.scalars(select(ProformaPilotageTramo)):
        hasta = form.get(f"pilot_hasta_{pt.id}")
        valor = form.get(f"pilot_valor_{pt.id}")
        pt.hasta_pies = float(hasta) if hasta not in (None, "") else None
        if valor not in (None, ""):
            try:
                pt.valor_usd = float(valor)
            except ValueError:
                pass

    for tt in db.scalars(select(ProformaTarifaTurno)):
        v = form.get(f"turno_{tt.id}")
        if v not in (None, ""):
            try:
                tt.valor_ars_dia = float(v)
            except ValueError:
                pass

    for cf in db.scalars(select(ProformaConceptoFijo)):
        valor = form.get(f"fijo_valor_{cf.id}")
        activo = form.get(f"fijo_activo_{cf.id}") == "on"
        if valor not in (None, ""):
            try:
                cf.valor_usd = float(valor)
            except ValueError:
                pass
        cf.activo = activo

    nombre_nuevo = (form.get("fijo_nuevo_nombre") or "").strip()
    if nombre_nuevo:
        try:
            valor_nuevo = float(form.get("fijo_nuevo_valor") or 0)
        except ValueError:
            valor_nuevo = 0.0
        ultimo = db.scalar(select(ProformaConceptoFijo).order_by(ProformaConceptoFijo.orden.desc()))
        orden = (ultimo.orden + 1) if ultimo else 0
        db.add(ProformaConceptoFijo(nombre=nombre_nuevo, valor_usd=valor_nuevo, activo=True, orden=orden))

    for bb in db.scalars(select(ProformaBunkerBoya)):
        for campo, clave in (("osro", "osro_usd"), ("boattrip", "boat_trip_usd"), ("boathora", "boat_hora_usd")):
            v = form.get(f"boya_{campo}_{bb.id}")
            if v not in (None, ""):
                try:
                    setattr(bb, clave, float(v))
                except ValueError:
                    pass
        bb.cobra_channel_anchor = form.get(f"boya_chanchor_{bb.id}") == "on"
        bb.cobra_pilotage = form.get(f"boya_pilotage_{bb.id}") == "on"
        bb.cobra_sipa = form.get(f"boya_sipa_{bb.id}") == "on"

    db.commit()
    return RedirectResponse(url="/proformador/config/formulas", status_code=302)
